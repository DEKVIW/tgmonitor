from __future__ import annotations

import hashlib
import io
import ipaddress
import json
import logging
import re
import socket
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta
from threading import Lock
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

from sqlalchemy import and_, case, exists, func, or_
from sqlalchemy.orm import Session

from app.models.models import (
    MediaWallCategory,
    MediaWallImportJob,
    MediaWallItem,
    MediaWallItemCategory,
    MediaWallLink,
    MediaWallSourceConfig,
    PanTransferSyncTask,
    ensure_runtime_storage_tables,
)
from app.services.link_check.parser import detect_platform_from_url, is_http_url, normalize_candidate_url
from app.services.resource_identity import normalize_match_key, parse_resource_identity


logger = logging.getLogger(__name__)

MEDIA_WALL_PUBLIC = "public"
MEDIA_WALL_HIDDEN = "hidden"
MEDIA_WALL_VISIBILITIES = {MEDIA_WALL_PUBLIC, MEDIA_WALL_HIDDEN}
MAX_IMPORT_ROWS = 5000
MAX_SOURCE_RESPONSE_BYTES = 2 * 1024 * 1024
METADATA_CACHE_TTL_SECONDS = 6 * 60 * 60
SOURCE_ENDPOINT_CACHE_TTL_SECONDS = 10 * 60
MEDIA_WALL_METADATA_NOT_SYNCED = "not_synced"
MEDIA_WALL_METADATA_PENDING = "pending"
MEDIA_WALL_METADATA_PROCESSING = "processing"
MEDIA_WALL_METADATA_READY = "ready"
MEDIA_WALL_METADATA_FAILED = "failed"

_metadata_cache: dict[str, tuple[float, dict[str, Any]]] = {}
_metadata_cache_lock = Lock()
_source_endpoint_cache: dict[str, tuple[float, list[str]]] = {}
_source_endpoint_cache_lock = Lock()


def _clean_text(value: Any, *, max_length: int = 255) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.replace("\u200b", " ").split())[:max_length].strip()


def _normalize_title(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", _clean_text(value, max_length=255).lower())


def _slugify(value: str) -> str:
    slug = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "-", _clean_text(value, max_length=120)).strip("-").lower()
    return slug or f"category-{hashlib.sha1(value.encode('utf-8')).hexdigest()[:10]}"


def _now() -> datetime:
    return datetime.utcnow()


def _fingerprint_url(url: str) -> str:
    normalized = _canonical_link_url(url)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _canonical_link_url(url: str) -> str:
    normalized = normalize_candidate_url(url)
    parsed = urlparse(normalized)
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, parsed.query, ""))


def _serialize_category(category: MediaWallCategory) -> dict[str, Any]:
    return {
        "id": int(category.id),
        "name": category.name,
        "slug": category.slug,
        "is_enabled": bool(category.is_enabled),
        "show_in_tabs": bool(category.is_enabled),
        "sort_order": int(category.sort_order or 0),
        "created_at": category.created_at,
        "updated_at": category.updated_at,
    }


def _serialize_link(link: MediaWallLink) -> dict[str, Any]:
    return {
        "id": int(link.id),
        "platform": link.platform,
        "url": link.url,
        "label": link.label,
        "is_primary": bool(link.is_primary),
        "is_enabled": bool(link.is_enabled),
        "source_type": link.source_type,
        "last_checked_at": link.last_checked_at,
        "last_status": link.last_status,
        "invalid_streak": int(link.invalid_streak or 0),
        "next_check_at": link.next_check_at,
        "last_error_message": link.last_error_message,
    }


def _serialize_link_health(links: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(links)
    enabled = sum(1 for link in links if link.get("is_enabled"))
    valid = sum(1 for link in links if str(link.get("last_status") or "unknown").lower() in {"valid", "healthy"})
    invalid = sum(1 for link in links if str(link.get("last_status") or "unknown").lower() == "invalid")
    pending = sum(1 for link in links if str(link.get("last_status") or "unknown").lower() in {"unknown", "checking"})
    uncertain = max(0, total - valid - invalid - pending)
    if not total:
        status = "no_links"
    elif valid:
        status = "valid"
    elif invalid == total:
        status = "invalid"
    elif uncertain:
        status = "uncertain"
    else:
        status = "pending"
    return {
        "total": total,
        "enabled": enabled,
        "valid": valid,
        "invalid": invalid,
        "pending": pending,
        "uncertain": uncertain,
        "status": status,
    }


def _extract_episode(value: Any) -> int | None:
    if isinstance(value, dict):
        for key in (
            "latest_episode",
            "current_episode",
            "episode",
            "synced_episode",
            "episode_number",
        ):
            candidate = value.get(key)
            try:
                number = int(candidate)
            except (TypeError, ValueError):
                number = 0
            if number > 0:
                return number
        for nested in value.values():
            result = _extract_episode(nested)
            if result is not None:
                return result
    elif isinstance(value, list):
        for nested in value:
            result = _extract_episode(nested)
            if result is not None:
                return result
    return None


def _follow_source_origin_title(task: PanTransferSyncTask) -> str:
    extra_json = dict(task.extra_json or {})
    source_origin = extra_json.get("source_origin")
    if not isinstance(source_origin, dict):
        return ""
    source_snapshot = source_origin.get("source_message_snapshot")
    if not isinstance(source_snapshot, dict):
        return ""
    return _clean_text(source_snapshot.get("title"), max_length=255)


def _follow_state_matches_source_origin(
    task: PanTransferSyncTask,
    state: Any,
) -> bool:
    if not isinstance(state, dict) or not state:
        return True
    origin_title = _follow_source_origin_title(task)
    state_title = _clean_text(state.get("resource_title") or state.get("core_title"), max_length=255)
    if not origin_title:
        # Keep legacy media-wall records visible until source identity is
        # explicitly available for a safe mismatch check.
        return True
    if not state_title:
        return False

    origin_identity = parse_resource_identity(origin_title)
    state_identity = parse_resource_identity(state_title)
    if not origin_identity.is_target_work or not state_identity.is_target_work:
        return normalize_match_key(origin_title) == normalize_match_key(state_title)
    origin_keys = {
        normalize_match_key(value)
        for value in (origin_identity.core_title, *origin_identity.aliases)
        if normalize_match_key(value)
    }
    state_key = normalize_match_key(state_identity.core_title)
    return bool(state_key and state_key in origin_keys)


def _is_trusted_current_share_identity(task: PanTransferSyncTask, state: Any) -> bool:
    if not isinstance(state, dict) or not state:
        return True
    return _follow_state_matches_source_origin(task, state)


def _get_trusted_follow_state(task: PanTransferSyncTask, key: str) -> dict[str, Any]:
    state = (task.extra_json or {}).get(key)
    if not isinstance(state, dict):
        return {}
    return dict(state) if _follow_state_matches_source_origin(task, state) else {}


def _get_trusted_follow_identity_snapshot(task: PanTransferSyncTask) -> dict[str, Any]:
    snapshot = (task.extra_json or {}).get("identity_snapshot")
    if not isinstance(snapshot, dict):
        return {}
    return dict(snapshot) if _follow_state_matches_source_origin(task, snapshot) else {}


def _get_trusted_follow_title(task: PanTransferSyncTask) -> str:
    origin_title = _follow_source_origin_title(task)
    candidates = (
        _get_trusted_follow_state(task, "applied_update_state"),
        _get_trusted_follow_state(task, "materialized_update_state"),
        _get_trusted_follow_identity_snapshot(task),
        (task.extra_json or {}).get("current_share_identity") or {},
    )
    for state in candidates:
        if not isinstance(state, dict) or not _follow_state_matches_source_origin(task, state):
            continue
        title = _clean_text(state.get("resource_title") or state.get("core_title"), max_length=255)
        if title:
            return title
    for candidate in (
        origin_title,
        task.work_title,
        (task.extra_json or {}).get("source_message_snapshot", {}).get("title")
        if isinstance((task.extra_json or {}).get("source_message_snapshot"), dict)
        else None,
        task.topic_title,
        task.task_name,
    ):
        title = _clean_text(candidate, max_length=255)
        if title and _follow_state_matches_source_origin(task, {"resource_title": title}):
            return title
    return ""


def _build_progress(task: PanTransferSyncTask | None) -> dict[str, Any]:
    if task is None:
        return {"episode": None, "label": "暂无追更进度", "task_state": None, "share_status": None}

    episode = None
    for state in (
        _get_trusted_follow_state(task, "applied_update_state"),
        _get_trusted_follow_state(task, "materialized_update_state"),
        _get_trusted_follow_identity_snapshot(task),
    ):
        episode = _extract_episode(state)
        if episode is not None:
            break
    if episode is None:
        identity = parse_resource_identity(_get_trusted_follow_title(task))
        episode = identity.episode
    return {
        "episode": episode,
        "label": f"更新至第 {episode} 集" if episode else "追更同步中",
        "task_state": task.task_state,
        "share_status": task.current_share_status,
        "last_checked_at": task.last_checked_at,
    }


def _parse_content_updated_at(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not isinstance(value, str):
        return None
    normalized = value.strip()
    if not normalized:
        return None
    if normalized.endswith("Z"):
        normalized = f"{normalized[:-1]}+00:00"
    try:
        return datetime.fromisoformat(normalized)
    except ValueError:
        return None


def _get_content_updated_at(task: PanTransferSyncTask | None) -> datetime | None:
    """Return only the actual update time for an active, unfinished follow task."""
    if task is None or str(task.status or "") != "active":
        return None
    for state in (
        _get_trusted_follow_state(task, "applied_update_state"),
        _get_trusted_follow_state(task, "materialized_update_state"),
    ):
        if str(state.get("is_completed") or "").lower() == "true":
            continue
        updated_at = _parse_content_updated_at(state.get("updated_at"))
        if updated_at is not None:
            return updated_at
    return None


def _public_item_ordering() -> tuple[Any, ...]:
    applied_update_state = func.jsonb_extract_path(PanTransferSyncTask.extra_json, "applied_update_state")
    content_updated_at = func.jsonb_extract_path_text(applied_update_state, "updated_at")
    is_completed = func.lower(func.jsonb_extract_path_text(applied_update_state, "is_completed"))
    is_active_follow = and_(
        PanTransferSyncTask.status == "active",
        or_(is_completed.is_(None), is_completed != "true"),
    )
    has_content_update = and_(is_active_follow, content_updated_at.isnot(None))
    return (
        case((has_content_update, 0), else_=1).asc(),
        case((has_content_update, content_updated_at), else_=None).desc().nullslast(),
        MediaWallItem.sort_order.asc(),
        MediaWallItem.id.desc(),
    )


def _validate_public_url(url: str) -> str:
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("只支持 http 或 https 地址")
    hostname = parsed.hostname.rstrip(".").lower()
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(hostname, parsed.port, type=socket.SOCK_STREAM)}
    except (OSError, ValueError) as exc:
        raise ValueError("无法解析采集源地址") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast or ip.is_reserved or ip.is_unspecified:
            raise ValueError("采集源不能访问内网或本机地址")
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", parsed.query, ""))


def _fetch_json(url: str, *, timeout_seconds: int) -> Any:
    safe_url = _validate_public_url(url)
    request = urllib.request.Request(
        safe_url,
        headers={"User-Agent": "TG-Monitor-MediaWall/1.0", "Accept": "application/json,text/plain;q=0.8"},
    )
    try:
        with urllib.request.urlopen(request, timeout=max(2, min(int(timeout_seconds), 20))) as response:
            payload = response.read(MAX_SOURCE_RESPONSE_BYTES + 1)
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        raise ValueError(f"采集源请求失败: {exc}") from exc
    if len(payload) > MAX_SOURCE_RESPONSE_BYTES:
        raise ValueError("采集源响应超过 2MB 限制")
    try:
        return json.loads(payload.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("采集源未返回有效 JSON") from exc


def _source_endpoints(source: MediaWallSourceConfig) -> list[str]:
    cache_key = ":".join(
        (
            str(source.id),
            str(source.source_type or ""),
            str(source.source_url or ""),
            source.updated_at.isoformat() if source.updated_at else "",
        )
    )
    now = time.monotonic()
    with _source_endpoint_cache_lock:
        cached = _source_endpoint_cache.get(cache_key)
        if cached and now - cached[0] < SOURCE_ENDPOINT_CACHE_TTL_SECONDS:
            return list(cached[1])

    if source.source_type in {"apple-cms", "api"}:
        endpoints = [_validate_public_url(source.source_url)]
        with _source_endpoint_cache_lock:
            _source_endpoint_cache[cache_key] = (now, endpoints)
        return endpoints

    payload = _fetch_json(source.source_url, timeout_seconds=source.timeout_seconds)
    if isinstance(payload, dict) and isinstance(payload.get("list"), list):
        endpoints = [_validate_public_url(source.source_url)]
        with _source_endpoint_cache_lock:
            _source_endpoint_cache[cache_key] = (now, endpoints)
        return endpoints
    entries: Any = []
    if isinstance(payload, dict):
        entries = payload.get("sources") or payload.get("sites") or payload.get("data") or []
    if not isinstance(entries, list):
        entries = []
    endpoints: list[str] = []
    for entry in entries:
        if isinstance(entry, str):
            candidate = entry
        elif isinstance(entry, dict):
            candidate = entry.get("api") or entry.get("api_url") or entry.get("url") or entry.get("ext") or ""
        else:
            candidate = ""
        if isinstance(candidate, str) and candidate.startswith(("http://", "https://")):
            try:
                endpoints.append(_validate_public_url(candidate))
            except ValueError:
                continue
    endpoints = endpoints[:20]
    with _source_endpoint_cache_lock:
        _source_endpoint_cache[cache_key] = (now, endpoints)
    return endpoints


def _query_endpoint(endpoint: str, title: str, *, timeout_seconds: int) -> dict[str, Any]:
    parsed = urlparse(endpoint)
    query = dict(parse_qsl(parsed.query, keep_blank_values=True))
    query.update({"ac": "videolist", "wd": title, "pg": "1"})
    request_url = urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, urlencode(query), ""))
    payload = _fetch_json(request_url, timeout_seconds=timeout_seconds)
    return payload if isinstance(payload, dict) else {}


def _parse_metadata(payload: dict[str, Any], query_title: str) -> dict[str, Any]:
    values: Any = payload.get("list")
    if isinstance(payload.get("data"), dict):
        values = payload["data"].get("list") or payload["data"].get("results") or values
    if not isinstance(values, list):
        return {}
    candidates = [item for item in values if isinstance(item, dict)]
    if not candidates:
        return {}
    normalized_query = _normalize_title(query_title)
    selected = next(
        (item for item in candidates if _normalize_title(item.get("vod_name") or item.get("name")) == normalized_query),
        candidates[0],
    )
    return {
        "title": _clean_text(selected.get("vod_name") or selected.get("name"), max_length=255),
        "poster_url": _clean_text(selected.get("vod_pic") or selected.get("pic"), max_length=1000),
        "type_name": _clean_text(selected.get("type_name") or selected.get("type"), max_length=120),
        "year": _clean_text(selected.get("vod_year") or selected.get("year"), max_length=16),
        "area": _clean_text(selected.get("vod_area") or selected.get("area"), max_length=80),
        "remarks": _clean_text(selected.get("vod_remarks") or selected.get("remarks"), max_length=120),
        "summary": _clean_text(selected.get("vod_content") or selected.get("content"), max_length=2000),
        "director": _clean_text(selected.get("vod_director") or selected.get("director"), max_length=255),
        "actor": _clean_text(selected.get("vod_actor") or selected.get("actor"), max_length=500),
    }


def _fetch_metadata(session: Session, query_title: str) -> dict[str, Any]:
    cache_key = _normalize_title(query_title)
    if not cache_key:
        return {}
    now = time.monotonic()
    with _metadata_cache_lock:
        cached = _metadata_cache.get(cache_key)
        if cached and now - cached[0] < METADATA_CACHE_TTL_SECONDS:
            return dict(cached[1])

    sources = (
        session.query(MediaWallSourceConfig)
        .filter(MediaWallSourceConfig.is_enabled.is_(True))
        .order_by(MediaWallSourceConfig.priority.asc(), MediaWallSourceConfig.id.asc())
        .all()
    )
    for source in sources:
        try:
            for endpoint in _source_endpoints(source):
                metadata = _parse_metadata(
                    _query_endpoint(endpoint, query_title, timeout_seconds=source.timeout_seconds),
                    query_title,
                )
                if metadata:
                    with _metadata_cache_lock:
                        _metadata_cache[cache_key] = (now, metadata)
                    return metadata
        except Exception as exc:
            logger.warning("media wall metadata source failed: source_id=%s error=%s", source.id, exc)
    with _metadata_cache_lock:
        _metadata_cache[cache_key] = (now, {})
    return {}


def _invalidate_metadata_cache(query_title: str | None = None) -> None:
    with _metadata_cache_lock:
        if query_title:
            _metadata_cache.pop(_normalize_title(query_title), None)
        else:
            _metadata_cache.clear()


def _invalidate_source_endpoint_cache(source_id: int | None = None) -> None:
    with _source_endpoint_cache_lock:
        if source_id is None:
            _source_endpoint_cache.clear()
            return
        prefix = f"{int(source_id)}:"
        for cache_key in list(_source_endpoint_cache):
            if cache_key.startswith(prefix):
                _source_endpoint_cache.pop(cache_key, None)


def _get_categories_for_item(session: Session, item_id: int, *, enabled_only: bool = False) -> list[MediaWallCategory]:
    query = (
        session.query(MediaWallCategory)
        .join(MediaWallItemCategory, MediaWallItemCategory.category_id == MediaWallCategory.id)
        .filter(MediaWallItemCategory.item_id == int(item_id))
        .order_by(MediaWallCategory.sort_order.asc(), MediaWallCategory.id.asc())
    )
    if enabled_only:
        query = query.filter(MediaWallCategory.is_enabled.is_(True))
    return query.all()


def _get_links_for_item(session: Session, item_id: int, *, enabled_only: bool = False) -> list[MediaWallLink]:
    query = session.query(MediaWallLink).filter(MediaWallLink.item_id == int(item_id))
    if enabled_only:
        query = query.filter(MediaWallLink.is_enabled.is_(True))
    return query.order_by(MediaWallLink.is_primary.desc(), MediaWallLink.id.asc()).all()


def _build_item_serialization_context(
    session: Session,
    items: Iterable[MediaWallItem],
    *,
    public: bool,
) -> tuple[dict[int, PanTransferSyncTask], dict[int, list[MediaWallLink]], dict[int, list[MediaWallCategory]]]:
    item_list = list(items)
    item_ids = [int(item.id) for item in item_list]
    if not item_ids:
        return {}, {}, {}

    sync_task_ids = {int(item.sync_task_id) for item in item_list if item.sync_task_id}
    tasks = (
        {
            int(task.id): task
            for task in session.query(PanTransferSyncTask).filter(PanTransferSyncTask.id.in_(sync_task_ids)).all()
        }
        if sync_task_ids
        else {}
    )
    links_by_item: dict[int, list[MediaWallLink]] = {}
    for link in (
        session.query(MediaWallLink)
        .filter(MediaWallLink.item_id.in_(item_ids))
        .order_by(MediaWallLink.item_id.asc(), MediaWallLink.is_primary.desc(), MediaWallLink.id.asc())
        .all()
    ):
        links_by_item.setdefault(int(link.item_id), []).append(link)

    category_query = (
        session.query(MediaWallItemCategory.item_id, MediaWallCategory)
        .join(MediaWallCategory, MediaWallItemCategory.category_id == MediaWallCategory.id)
        .filter(MediaWallItemCategory.item_id.in_(item_ids))
    )
    if public:
        category_query = category_query.filter(MediaWallCategory.is_enabled.is_(True))
    categories_by_item: dict[int, list[MediaWallCategory]] = {}
    for item_id, category in category_query.order_by(
        MediaWallItemCategory.item_id.asc(), MediaWallCategory.sort_order.asc(), MediaWallCategory.id.asc()
    ).all():
        categories_by_item.setdefault(int(item_id), []).append(category)
    return tasks, links_by_item, categories_by_item


def _serialize_item(
    session: Session,
    item: MediaWallItem,
    *,
    public: bool,
    include_metadata: bool,
    tasks: dict[int, PanTransferSyncTask] | None = None,
    links_by_item: dict[int, list[MediaWallLink]] | None = None,
    categories_by_item: dict[int, list[MediaWallCategory]] | None = None,
) -> dict[str, Any]:
    if tasks is None:
        task = session.get(PanTransferSyncTask, item.sync_task_id) if item.sync_task_id else None
    else:
        task = tasks.get(int(item.sync_task_id)) if item.sync_task_id else None
    stored_metadata = item.metadata_json if isinstance(item.metadata_json, dict) else {}
    metadata = dict(stored_metadata) if include_metadata else {}
    links = links_by_item.get(int(item.id), []) if links_by_item is not None else _get_links_for_item(session, item.id, enabled_only=False)
    categories = categories_by_item.get(int(item.id), []) if categories_by_item is not None else _get_categories_for_item(session, item.id, enabled_only=public)
    poster_url = item.custom_poster_url or metadata.get("poster_url") or None
    progress = _build_progress(task)
    content_updated_at = _get_content_updated_at(task)
    serialized_links: list[dict[str, Any]] = []
    raw_current_share_url = _canonical_link_url(task.current_share_url) if task and task.current_share_url else ""
    current_share_identity = (task.extra_json or {}).get("current_share_identity") if task else None
    current_share_identity_is_trusted = bool(
        not public
        or _is_trusted_current_share_identity(task, current_share_identity)
    )
    blocked_current_share_url = raw_current_share_url if not current_share_identity_is_trusted else ""
    current_share_url = raw_current_share_url if current_share_identity_is_trusted else ""
    task_is_active = bool(task and str(task.status or "") == "active")
    task_share_status = str(task.current_share_status or "unknown").lower() if task else "unknown"
    task_share_enabled = not task_is_active or task_share_status != "invalid"
    for link in links:
        serialized_link = _serialize_link(link)
        if public and blocked_current_share_url and serialized_link["url"] == blocked_current_share_url:
            continue
        if current_share_url and serialized_link["url"] == current_share_url and task and task.status != "deleted":
            serialized_link.update(
                {
                    "source_type": "sync",
                    "last_checked_at": task.last_checked_at,
                    "last_status": task_share_status if task_is_active else "unknown",
                    "is_enabled": task_share_enabled,
                    "invalid_streak": 0,
                    "next_check_at": None,
                    "last_error_message": None,
                }
            )
        if public and (not serialized_link["is_enabled"] or serialized_link["last_status"] == "invalid"):
            continue
        serialized_links.append(serialized_link)
    if (
        current_share_url
        and is_http_url(current_share_url)
        and task
        and task.status != "deleted"
        and all(link["url"] != current_share_url for link in serialized_links)
        and (not public or task_share_enabled)
    ):
        serialized_links.insert(
            0,
            {
                "id": -int(task.id),
                "platform": detect_platform_from_url(current_share_url),
                "url": current_share_url,
                "label": "追更同步当前链接",
                "is_primary": not any(link["is_primary"] for link in serialized_links),
                "is_enabled": True,
                "source_type": "sync",
                "last_checked_at": task.last_checked_at,
                "last_status": task_share_status if task_is_active else "unknown",
                "invalid_streak": 0,
                "next_check_at": None,
                "last_error_message": None,
            },
        )
    return {
        "id": int(item.id),
        "display_title": item.display_title,
        "metadata_query_title": item.metadata_query_title if not public else None,
        "visibility": item.visibility,
        "sort_order": int(item.sort_order or 0),
        "poster_url": poster_url,
        "type_name": metadata.get("type_name"),
        "year": metadata.get("year"),
        "area": metadata.get("area"),
        "remarks": metadata.get("remarks"),
        "summary": metadata.get("summary"),
        "director": metadata.get("director"),
        "actor": metadata.get("actor"),
        "progress": progress,
        "links": serialized_links,
        "link_health": _serialize_link_health(serialized_links),
        "link_count": len(serialized_links),
        "categories": [_serialize_category(category) for category in categories],
        "sync_task_id": int(item.sync_task_id) if item.sync_task_id else None,
        "metadata_status": str(item.metadata_status or MEDIA_WALL_METADATA_NOT_SYNCED),
        "metadata_last_synced_at": item.metadata_last_synced_at,
        "metadata_last_error": item.metadata_last_error if not public else None,
        "last_content_updated_at": content_updated_at,
        "created_at": item.created_at,
        "updated_at": item.updated_at,
    }


def list_public_categories(session: Session) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    categories = (
        session.query(MediaWallCategory)
        .filter(MediaWallCategory.is_enabled.is_(True))
        .order_by(MediaWallCategory.sort_order.asc(), MediaWallCategory.id.asc())
        .all()
    )
    return [{"id": 0, "name": "全部", "slug": "all", "is_enabled": True, "show_in_tabs": True, "sort_order": -1}] + [
        _serialize_category(category) for category in categories
    ]


def _public_link_exists_filter() -> Any:
    current_share_invalid = exists().where(
        PanTransferSyncTask.id == MediaWallItem.sync_task_id,
        PanTransferSyncTask.status == "active",
        PanTransferSyncTask.current_share_status == "invalid",
        PanTransferSyncTask.current_share_url == MediaWallLink.url,
    ).correlate(MediaWallItem, MediaWallLink)
    return or_(
        exists().where(
            MediaWallLink.item_id == MediaWallItem.id,
            MediaWallLink.is_enabled.is_(True),
            MediaWallLink.last_status != "invalid",
            ~current_share_invalid,
        ).correlate(MediaWallItem),
        exists().where(
            PanTransferSyncTask.id == MediaWallItem.sync_task_id,
            PanTransferSyncTask.status != "deleted",
            PanTransferSyncTask.current_share_url.isnot(None),
            or_(
                PanTransferSyncTask.status != "active",
                PanTransferSyncTask.current_share_status != "invalid",
            ),
        ).correlate(MediaWallItem),
    )


def list_public_items(
    session: Session,
    *,
    category_slug: str | None = None,
    search: str | None = None,
    page: int = 1,
    page_size: int = 24,
) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    query = session.query(MediaWallItem).filter(MediaWallItem.visibility == MEDIA_WALL_PUBLIC)
    normalized_search = _clean_text(search, max_length=120)
    if normalized_search:
        query = query.filter(MediaWallItem.display_title.ilike(f"%{normalized_search}%"))
    normalized_category = _clean_text(category_slug, max_length=160).lower()
    if normalized_category and normalized_category != "all":
        category = (
            session.query(MediaWallCategory)
            .filter(
                func.lower(MediaWallCategory.slug) == normalized_category,
                MediaWallCategory.is_enabled.is_(True),
            )
            .first()
        )
        if category is None:
            return {"items": [], "total": 0, "page": max(1, page), "page_size": min(max(1, page_size), 60)}
        query = query.join(MediaWallItemCategory, MediaWallItemCategory.item_id == MediaWallItem.id).filter(
            MediaWallItemCategory.category_id == category.id
        )
    query = query.filter(_public_link_exists_filter())
    safe_page = max(1, int(page or 1))
    safe_size = min(max(1, int(page_size or 24)), 60)
    candidate_items = (
        query.outerjoin(PanTransferSyncTask, PanTransferSyncTask.id == MediaWallItem.sync_task_id)
        .all()
    )
    serialization_context = _build_item_serialization_context(session, candidate_items, public=True)
    serialized_items = [
        _serialize_item(
            session,
            item,
            public=True,
            include_metadata=True,
            tasks=serialization_context[0],
            links_by_item=serialization_context[1],
            categories_by_item=serialization_context[2],
        )
        for item in candidate_items
    ]
    serialized_items = [item for item in serialized_items if item["links"]]

    def public_sort_key(item: dict[str, Any]) -> tuple[int, float, int, int]:
        updated_at = _parse_content_updated_at(item.get("last_content_updated_at"))
        if updated_at is None:
            return (1, 0.0, int(item["sort_order"]), -int(item["id"]))
        return (0, -updated_at.timestamp(), int(item["sort_order"]), -int(item["id"]))

    serialized_items.sort(key=public_sort_key)
    start = (safe_page - 1) * safe_size
    return {
        "items": serialized_items[start : start + safe_size],
        "total": len(serialized_items),
        "page": safe_page,
        "page_size": safe_size,
    }


def queue_media_wall_metadata_sync(session: Session) -> dict[str, int | str]:
    ensure_runtime_storage_tables()
    items = session.query(MediaWallItem).order_by(MediaWallItem.id.asc()).all()
    queued_count = 0
    skipped_count = 0
    for item in items:
        current_status = str(item.metadata_status or MEDIA_WALL_METADATA_NOT_SYNCED)
        if current_status == MEDIA_WALL_METADATA_READY:
            skipped_count += 1
            continue
        if current_status in {MEDIA_WALL_METADATA_PENDING, MEDIA_WALL_METADATA_PROCESSING}:
            skipped_count += 1
            continue
        _invalidate_metadata_cache(item.metadata_query_title)
        item.metadata_status = MEDIA_WALL_METADATA_PENDING
        item.metadata_last_error = None
        queued_count += 1
    session.flush()
    return {
        "status": "queued" if queued_count else "up_to_date",
        "queued_count": queued_count,
        "skipped_count": skipped_count,
    }


def process_next_media_wall_metadata_sync(session: Session, *, worker_name: str) -> bool:
    ensure_runtime_storage_tables()
    item = (
        session.query(MediaWallItem)
        .filter(MediaWallItem.metadata_status == MEDIA_WALL_METADATA_PENDING)
        .order_by(MediaWallItem.id.asc())
        .with_for_update(skip_locked=True)
        .first()
    )
    if item is None:
        return False

    item.metadata_status = MEDIA_WALL_METADATA_PROCESSING
    item.metadata_last_error = None
    session.flush()
    try:
        metadata = _fetch_metadata(session, item.metadata_query_title)
        if not metadata:
            item.metadata_status = MEDIA_WALL_METADATA_FAILED
            item.metadata_last_error = "未从影视采集源获取到资料"
        else:
            item.metadata_json = metadata
            item.metadata_status = MEDIA_WALL_METADATA_READY
            item.metadata_last_synced_at = _now()
            item.metadata_last_error = None
    except Exception as exc:
        item.metadata_status = MEDIA_WALL_METADATA_FAILED
        item.metadata_last_error = _clean_text(exc, max_length=2000) or type(exc).__name__
        logger.exception("media wall metadata sync failed: item_id=%s", item.id)
    item.updated_by = worker_name
    session.flush()
    return True


def get_public_item(session: Session, item_id: int) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    item = session.query(MediaWallItem).filter(MediaWallItem.id == int(item_id), MediaWallItem.visibility == MEDIA_WALL_PUBLIC).first()
    if item is None:
        raise LookupError("影视条目不存在")
    serialized = _serialize_item(session, item, public=True, include_metadata=True)
    if not serialized["links"]:
        raise LookupError("影视条目暂无有效网盘入口")
    return serialized


def list_admin_categories(session: Session) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    return [
        _serialize_category(category)
        for category in session.query(MediaWallCategory).order_by(MediaWallCategory.sort_order.asc(), MediaWallCategory.id.asc()).all()
    ]


def create_category(session: Session, *, name: str, slug: str | None, sort_order: int, is_enabled: bool, operator: str) -> dict[str, Any]:
    clean_name = _clean_text(name, max_length=120)
    if not clean_name:
        raise ValueError("分类名称不能为空")
    clean_slug = _slugify(slug or clean_name)
    if session.query(MediaWallCategory).filter(func.lower(MediaWallCategory.slug) == clean_slug.lower()).first():
        raise ValueError("分类路径已存在")
    category = MediaWallCategory(
        name=clean_name,
        slug=clean_slug,
        sort_order=int(sort_order or 0),
        is_enabled=bool(is_enabled),
        show_in_tabs=bool(is_enabled),
        created_by=operator,
        updated_by=operator,
    )
    session.add(category)
    session.flush()
    return _serialize_category(category)


def update_category(session: Session, *, category_id: int, name: str, slug: str | None, sort_order: int, is_enabled: bool, operator: str) -> dict[str, Any]:
    category = session.get(MediaWallCategory, int(category_id))
    if category is None:
        raise LookupError("分类不存在")
    clean_name = _clean_text(name, max_length=120)
    if not clean_name:
        raise ValueError("分类名称不能为空")
    clean_slug = _slugify(slug or clean_name)
    conflict = (
        session.query(MediaWallCategory)
        .filter(func.lower(MediaWallCategory.slug) == clean_slug.lower(), MediaWallCategory.id != category.id)
        .first()
    )
    if conflict:
        raise ValueError("分类路径已存在")
    category.name = clean_name
    category.slug = clean_slug
    category.sort_order = int(sort_order or 0)
    category.is_enabled = bool(is_enabled)
    category.show_in_tabs = bool(is_enabled)
    category.updated_by = operator
    session.flush()
    return _serialize_category(category)


def disable_category(session: Session, *, category_id: int) -> dict[str, Any]:
    category = session.get(MediaWallCategory, int(category_id))
    if category is None:
        raise LookupError("分类不存在")
    category.is_enabled = False
    category.show_in_tabs = False
    session.flush()
    return _serialize_category(category)


def list_admin_items(
    session: Session,
    *,
    search: str | None = None,
    visibility: str | None = None,
    category_id: int | None = None,
) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    query = session.query(MediaWallItem)
    if search:
        query = query.filter(MediaWallItem.display_title.ilike(f"%{_clean_text(search, max_length=120)}%"))
    if visibility in MEDIA_WALL_VISIBILITIES:
        query = query.filter(MediaWallItem.visibility == visibility)
    if category_id is not None and int(category_id) > 0:
        query = query.join(MediaWallItemCategory, MediaWallItemCategory.item_id == MediaWallItem.id).filter(
            MediaWallItemCategory.category_id == int(category_id)
        )
    items = query.order_by(MediaWallItem.sort_order.asc(), MediaWallItem.updated_at.desc(), MediaWallItem.id.desc()).all()
    serialization_context = _build_item_serialization_context(session, items, public=False)
    return [
        _serialize_item(
            session,
            item,
            public=False,
            include_metadata=False,
            tasks=serialization_context[0],
            links_by_item=serialization_context[1],
            categories_by_item=serialization_context[2],
        )
        for item in items
    ]


def _delete_item_rows(session: Session, item_ids: Iterable[int]) -> list[MediaWallItem]:
    normalized_ids = sorted({int(item_id) for item_id in item_ids if int(item_id) > 0})
    if not normalized_ids:
        return []
    rows = session.query(MediaWallItem).filter(MediaWallItem.id.in_(normalized_ids)).all()
    found_ids = {int(row.id) for row in rows}
    missing_ids = [item_id for item_id in normalized_ids if item_id not in found_ids]
    if missing_ids:
        raise LookupError(f"影视条目不存在: {missing_ids[0]}")
    session.query(MediaWallItemCategory).filter(MediaWallItemCategory.item_id.in_(normalized_ids)).delete(
        synchronize_session=False
    )
    session.query(MediaWallLink).filter(MediaWallLink.item_id.in_(normalized_ids)).delete(synchronize_session=False)
    for row in rows:
        session.delete(row)
    session.flush()
    return rows


def delete_item(session: Session, *, item_id: int) -> dict[str, Any]:
    row = session.get(MediaWallItem, int(item_id))
    if row is None:
        raise LookupError("影视条目不存在")
    result = _serialize_item(session, row, public=False, include_metadata=False)
    _delete_item_rows(session, [row.id])
    return result


def delete_items(session: Session, *, item_ids: Iterable[int]) -> list[int]:
    rows = _delete_item_rows(session, item_ids)
    return [int(row.id) for row in rows]


def delete_items_for_sync_task(session: Session, *, sync_task_id: int) -> int:
    item_ids = [
        int(item_id)
        for (item_id,) in session.query(MediaWallItem.id)
        .filter(MediaWallItem.sync_task_id == int(sync_task_id))
        .all()
    ]
    return len(delete_items(session, item_ids=item_ids))


def _validate_category_ids(session: Session, category_ids: Iterable[int]) -> list[int]:
    ids = sorted({int(value) for value in category_ids if int(value) > 0})
    if not ids:
        return []
    found = {int(row.id) for row in session.query(MediaWallCategory).filter(MediaWallCategory.id.in_(ids)).all()}
    if found != set(ids):
        raise ValueError("存在无效的影视墙分类")
    return ids


def _replace_item_categories(session: Session, *, item_id: int, category_ids: Iterable[int], primary_category_id: int | None = None) -> None:
    ids = _validate_category_ids(session, category_ids)
    session.query(MediaWallItemCategory).filter(MediaWallItemCategory.item_id == int(item_id)).delete(synchronize_session=False)
    primary = int(primary_category_id) if primary_category_id in ids else (ids[0] if ids else None)
    for category_id in ids:
        session.add(MediaWallItemCategory(item_id=int(item_id), category_id=category_id, is_primary=category_id == primary))


def _replace_item_links(session: Session, *, item_id: int, link_payloads: Iterable[dict[str, Any]]) -> None:
    item = session.get(MediaWallItem, int(item_id))
    if item is None:
        raise LookupError("影视条目不存在")

    normalized_payloads: list[dict[str, Any]] = []
    seen_fingerprints: set[str] = set()
    for payload in link_payloads:
        clean_url = _canonical_link_url(str(payload.get("url") or ""))
        if not is_http_url(clean_url):
            raise ValueError("网盘链接格式无效")
        fingerprint = _fingerprint_url(clean_url)
        if fingerprint in seen_fingerprints:
            continue
        seen_fingerprints.add(fingerprint)
        normalized_payloads.append(
            {
                "id": int(payload["id"]) if payload.get("id") is not None and int(payload["id"]) > 0 else None,
                "url": clean_url,
                "fingerprint": fingerprint,
                "label": _clean_text(payload.get("label"), max_length=120) or None,
                "source_type": str(payload.get("source_type") or "manual"),
                "is_enabled": bool(payload.get("is_enabled", True)),
                "is_primary": bool(payload.get("is_primary", False)),
            }
        )

    existing_links = _get_links_for_item(session, int(item_id), enabled_only=False)
    by_fingerprint = {str(link.url_fingerprint): link for link in existing_links}
    by_id = {int(link.id): link for link in existing_links}
    task = session.get(PanTransferSyncTask, item.sync_task_id) if item.sync_task_id else None
    protected_url = ""
    if task and task.status != "deleted" and task.current_share_url:
        protected_url = _canonical_link_url(task.current_share_url)
    requested_fingerprints: set[str] = set()
    requested_links: list[MediaWallLink] = []

    for payload in normalized_payloads:
        requested_fingerprints.add(payload["fingerprint"])
        if protected_url and payload["url"] == protected_url:
            protected_link = by_fingerprint.get(payload["fingerprint"])
            if protected_link is not None:
                requested_links.append(protected_link)
            continue

        link = by_id.get(payload["id"]) if payload["id"] is not None else None
        if link is not None and str(link.url_fingerprint) != payload["fingerprint"]:
            link = None
        if link is None:
            link = by_fingerprint.get(payload["fingerprint"])
        if link is None:
            link = MediaWallLink(
                item_id=int(item_id),
                platform=detect_platform_from_url(payload["url"]),
                url=payload["url"],
                url_fingerprint=payload["fingerprint"],
                label=payload["label"],
                is_primary=payload["is_primary"],
                is_enabled=payload["is_enabled"],
                source_type=payload["source_type"] if payload["source_type"] in {"manual", "import"} else "manual",
                last_status="unknown",
            )
            session.add(link)
            session.flush()
            by_fingerprint[payload["fingerprint"]] = link
        else:
            link.label = payload["label"]
            link.is_enabled = payload["is_enabled"]
            if link.source_type != "sync":
                link.is_primary = payload["is_primary"]
        requested_links.append(link)

    for link in existing_links:
        if link.source_type == "sync":
            if protected_url and link.url == protected_url:
                continue
            session.delete(link)
            continue
        if protected_url and link.url == protected_url:
            continue
        if str(link.url_fingerprint) not in requested_fingerprints:
            session.delete(link)

    primary_link = next((link for link in requested_links if link.is_primary), None)
    if primary_link is None:
        primary_link = next((link for link in requested_links if link.is_enabled), None)
    if primary_link is not None:
        session.query(MediaWallLink).filter(
            MediaWallLink.item_id == int(item_id),
            MediaWallLink.id != primary_link.id,
        ).update({MediaWallLink.is_primary: False}, synchronize_session=False)
        primary_link.is_primary = True


def _upsert_link(session: Session, *, item_id: int, url: str, label: str | None, source_type: str, is_primary: bool = False) -> MediaWallLink | None:
    clean_url = _canonical_link_url(url)
    if not is_http_url(clean_url):
        return None
    fingerprint = _fingerprint_url(clean_url)
    link = (
        session.query(MediaWallLink)
        .filter(MediaWallLink.item_id == int(item_id), MediaWallLink.url_fingerprint == fingerprint)
        .first()
    )
    if link is None:
        link = MediaWallLink(
            item_id=int(item_id),
            platform=detect_platform_from_url(clean_url),
            url=clean_url,
            url_fingerprint=fingerprint,
            label=_clean_text(label, max_length=120) or None,
            is_primary=bool(is_primary),
            source_type=source_type if source_type in {"manual", "import"} else "manual",
        )
        session.add(link)
    else:
        link.is_enabled = True
        link.invalid_streak = 0
        link.next_check_at = None
        link.last_checked_at = None
        link.last_status = "unknown"
        link.last_error_message = None
        if label:
            link.label = _clean_text(label, max_length=120)
        if link.source_type != "sync" and source_type in {"manual", "import"}:
            link.source_type = source_type
        if is_primary:
            link.is_primary = True
    session.flush()
    if is_primary:
        session.query(MediaWallLink).filter(MediaWallLink.item_id == int(item_id), MediaWallLink.id != link.id).update(
            {MediaWallLink.is_primary: False}, synchronize_session=False
        )
    return link


def create_item(session: Session, *, display_title: str, metadata_query_title: str | None, visibility: str, sort_order: int, category_ids: list[int], sync_task_id: int | None, custom_poster_url: str | None, links: list[dict[str, Any]], operator: str) -> dict[str, Any]:
    title = _clean_text(display_title, max_length=255)
    if not title:
        raise ValueError("影视名称不能为空")
    if visibility not in MEDIA_WALL_VISIBILITIES:
        raise ValueError("无效的影视公开状态")
    if sync_task_id is not None and session.get(PanTransferSyncTask, int(sync_task_id)) is None:
        raise ValueError("追更任务不存在")
    clean_metadata_query_title = _clean_text(metadata_query_title or title, max_length=255) or title
    _invalidate_metadata_cache(clean_metadata_query_title)
    item = MediaWallItem(
        display_title=title,
        metadata_query_title=clean_metadata_query_title,
        visibility=visibility,
        sort_order=int(sort_order or 0),
        sync_task_id=int(sync_task_id) if sync_task_id else None,
        custom_poster_url=_clean_text(custom_poster_url, max_length=1000) or None,
        metadata_status=MEDIA_WALL_METADATA_PENDING,
        created_by=operator,
        updated_by=operator,
    )
    session.add(item)
    session.flush()
    _replace_item_categories(session, item_id=item.id, category_ids=category_ids)
    for index, link_data in enumerate(links):
        _upsert_link(
            session,
            item_id=item.id,
            url=str(link_data.get("url") or ""),
            label=link_data.get("label"),
            source_type=str(link_data.get("source_type") or "manual"),
            is_primary=bool(link_data.get("is_primary") or index == 0),
        )
    session.flush()
    return _serialize_item(session, item, public=False, include_metadata=False)


def update_item(session: Session, *, item_id: int, display_title: str, metadata_query_title: str | None, visibility: str, sort_order: int, category_ids: list[int], sync_task_id: int | None, custom_poster_url: str | None, links: list[dict[str, Any]] | None, operator: str) -> dict[str, Any]:
    item = session.get(MediaWallItem, int(item_id))
    if item is None:
        raise LookupError("影视条目不存在")
    title = _clean_text(display_title, max_length=255)
    if not title:
        raise ValueError("影视名称不能为空")
    if visibility not in MEDIA_WALL_VISIBILITIES:
        raise ValueError("无效的影视公开状态")
    if sync_task_id is not None and session.get(PanTransferSyncTask, int(sync_task_id)) is None:
        raise ValueError("追更任务不存在")
    next_metadata_query_title = _clean_text(metadata_query_title or title, max_length=255) or title
    metadata_query_changed = item.metadata_query_title != next_metadata_query_title
    previous_metadata_query_title = item.metadata_query_title
    item.display_title = title
    item.metadata_query_title = next_metadata_query_title
    item.visibility = visibility
    item.sort_order = int(sort_order or 0)
    item.sync_task_id = int(sync_task_id) if sync_task_id else None
    item.custom_poster_url = _clean_text(custom_poster_url, max_length=1000) or None
    if metadata_query_changed:
        _invalidate_metadata_cache(previous_metadata_query_title)
        _invalidate_metadata_cache(next_metadata_query_title)
        item.metadata_json = {}
        item.metadata_last_synced_at = None
        item.metadata_status = MEDIA_WALL_METADATA_PENDING
        item.metadata_last_error = None
    elif str(item.metadata_status or MEDIA_WALL_METADATA_NOT_SYNCED) in {MEDIA_WALL_METADATA_NOT_SYNCED, MEDIA_WALL_METADATA_FAILED}:
        _invalidate_metadata_cache(next_metadata_query_title)
        item.metadata_status = MEDIA_WALL_METADATA_PENDING
        item.metadata_last_error = None
    item.updated_by = operator
    _replace_item_categories(session, item_id=item.id, category_ids=category_ids)
    if links is not None:
        _replace_item_links(session, item_id=item.id, link_payloads=links)
    session.flush()
    return _serialize_item(session, item, public=False, include_metadata=False)


def add_item_link(session: Session, *, item_id: int, url: str, label: str | None, source_type: str, is_primary: bool) -> dict[str, Any]:
    if session.get(MediaWallItem, int(item_id)) is None:
        raise LookupError("影视条目不存在")
    if source_type == "sync":
        raise ValueError("追更同步链接由追更任务自动管理")
    link = _upsert_link(session, item_id=item_id, url=url, label=label, source_type=source_type, is_primary=is_primary)
    if link is None:
        raise ValueError("网盘链接格式无效")
    return _serialize_link(link)


def update_item_link(session: Session, *, link_id: int, url: str, label: str | None, is_enabled: bool, is_primary: bool) -> dict[str, Any]:
    link = session.get(MediaWallLink, int(link_id))
    if link is None:
        raise LookupError("网盘链接不存在")
    if link.source_type == "sync":
        raise ValueError("追更同步链接由追更任务自动管理")
    clean_url = _canonical_link_url(url)
    if not is_http_url(clean_url):
        raise ValueError("网盘链接格式无效")
    new_fingerprint = _fingerprint_url(clean_url)
    conflict = session.query(MediaWallLink).filter(
        MediaWallLink.item_id == link.item_id,
        MediaWallLink.url_fingerprint == new_fingerprint,
        MediaWallLink.id != link.id,
    ).first()
    if conflict is not None:
        raise ValueError("该影视条目已经存在相同链接")
    url_changed = link.url_fingerprint != new_fingerprint
    link.url = clean_url
    link.platform = detect_platform_from_url(clean_url)
    link.url_fingerprint = new_fingerprint
    link.label = _clean_text(label, max_length=120) or None
    link.is_enabled = bool(is_enabled)
    link.is_primary = bool(is_primary)
    if url_changed:
        link.last_checked_at = None
        link.last_status = "unknown"
        link.invalid_streak = 0
        link.next_check_at = None
        link.last_error_message = None
    if is_primary:
        session.query(MediaWallLink).filter(MediaWallLink.item_id == link.item_id, MediaWallLink.id != link.id).update(
            {MediaWallLink.is_primary: False}, synchronize_session=False
        )
    session.flush()
    return _serialize_link(link)


def delete_item_link(session: Session, *, link_id: int) -> dict[str, Any]:
    link = session.get(MediaWallLink, int(link_id))
    if link is None:
        raise LookupError("网盘链接不存在")
    if link.source_type == "sync":
        raise ValueError("追更同步链接由追更任务自动管理")
    result = _serialize_link(link)
    session.delete(link)
    session.flush()
    return result


def list_sync_tasks(session: Session) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    rows = (
        session.query(PanTransferSyncTask)
        .filter(PanTransferSyncTask.status != "deleted")
        .order_by(PanTransferSyncTask.updated_at.desc(), PanTransferSyncTask.id.desc())
        .limit(200)
        .all()
    )
    return [
        {
            "id": int(row.id),
            "task_name": row.task_name,
            "work_title": row.work_title,
            "topic_title": row.topic_title,
            "status": row.status,
            "task_state": row.task_state,
            "current_share_url": row.current_share_url,
            "current_share_status": row.current_share_status,
            "progress": _build_progress(row),
        }
        for row in rows
    ]


def list_sources(session: Session) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    rows = session.query(MediaWallSourceConfig).order_by(MediaWallSourceConfig.priority.asc(), MediaWallSourceConfig.id.asc()).all()
    return [_serialize_source(row) for row in rows]


def _serialize_source(source: MediaWallSourceConfig) -> dict[str, Any]:
    return {
        "id": int(source.id),
        "name": source.name,
        "source_url": source.source_url,
        "source_type": source.source_type,
        "is_enabled": bool(source.is_enabled),
        "priority": int(source.priority or 100),
        "timeout_seconds": int(source.timeout_seconds or 8),
        "last_checked_at": source.last_checked_at,
        "last_refreshed_at": source.last_refreshed_at,
        "last_error_message": source.last_error_message,
        "endpoint_count": (source.extra_json or {}).get("endpoint_count", 0),
    }


def create_source(session: Session, *, name: str, source_url: str, source_type: str, is_enabled: bool, priority: int, timeout_seconds: int) -> dict[str, Any]:
    clean_url = _validate_public_url(source_url)
    _invalidate_source_endpoint_cache()
    _invalidate_metadata_cache()
    source = MediaWallSourceConfig(
        name=_clean_text(name, max_length=120) or "影视采集源",
        source_url=clean_url,
        source_type=_clean_text(source_type, max_length=32) or "source-list",
        is_enabled=bool(is_enabled),
        priority=max(0, int(priority)),
        timeout_seconds=min(20, max(2, int(timeout_seconds))),
    )
    session.add(source)
    session.flush()
    return _serialize_source(source)


def update_source(session: Session, *, source_id: int, name: str, source_url: str, source_type: str, is_enabled: bool, priority: int, timeout_seconds: int) -> dict[str, Any]:
    source = session.get(MediaWallSourceConfig, int(source_id))
    if source is None:
        raise LookupError("采集源不存在")
    _invalidate_source_endpoint_cache(source.id)
    _invalidate_metadata_cache()
    source.name = _clean_text(name, max_length=120) or "影视采集源"
    source.source_url = _validate_public_url(source_url)
    source.source_type = _clean_text(source_type, max_length=32) or "source-list"
    source.is_enabled = bool(is_enabled)
    source.priority = max(0, int(priority))
    source.timeout_seconds = min(20, max(2, int(timeout_seconds)))
    session.flush()
    return _serialize_source(source)


def test_source(session: Session, *, source_id: int, refresh: bool = False) -> dict[str, Any]:
    source = session.get(MediaWallSourceConfig, int(source_id))
    if source is None:
        raise LookupError("采集源不存在")
    source.last_checked_at = _now()
    try:
        endpoints = _source_endpoints(source)
        if not endpoints:
            raise ValueError("采集源未发现可用 API 地址")
        source.last_error_message = None
        source.extra_json = {**(source.extra_json or {}), "endpoint_count": len(endpoints)}
        if refresh:
            source.last_refreshed_at = _now()
        session.flush()
        return {"source": _serialize_source(source), "ok": True, "endpoint_count": len(endpoints)}
    except Exception as exc:
        source.last_error_message = _clean_text(exc, max_length=1000)
        session.flush()
        return {"source": _serialize_source(source), "ok": False, "endpoint_count": 0, "error": source.last_error_message}


def _split_import_row(row: list[Any]) -> tuple[str, list[str]]:
    values = [_clean_text(value, max_length=2000) for value in row]
    values = [value for value in values if value]
    if not values:
        return "", []
    return values[0], values[1:]


def _build_import_preview(rows: Iterable[list[Any]]) -> dict[str, Any]:
    records: list[dict[str, Any]] = []
    invalid_rows: list[dict[str, Any]] = []
    seen_titles: set[str] = set()
    for index, row in enumerate(rows, start=1):
        if index > MAX_IMPORT_ROWS:
            invalid_rows.append({"row": index, "reason": f"超过 {MAX_IMPORT_ROWS} 行限制"})
            break
        title, raw_links = _split_import_row(row)
        if not title and not raw_links:
            continue
        if index == 1 and title.lower() in {"影视名称", "名称", "title", "name"}:
            continue
        valid_links: list[dict[str, Any]] = []
        for raw_link in raw_links:
            normalized = normalize_candidate_url(raw_link)
            if is_http_url(normalized):
                valid_links.append({"url": normalized, "platform": detect_platform_from_url(normalized)})
        if not title or not valid_links:
            invalid_rows.append({"row": index, "title": title, "reason": "影视名称或有效网盘链接为空"})
            continue
        normalized_title = _normalize_title(title)
        if normalized_title in seen_titles:
            existing = next(record for record in records if _normalize_title(record["display_title"]) == normalized_title)
            existing["links"].extend(valid_links)
            continue
        seen_titles.add(normalized_title)
        records.append({"display_title": title, "links": valid_links})
    for record in records:
        unique_links: dict[str, dict[str, Any]] = {}
        for link in record["links"]:
            unique_links[_fingerprint_url(link["url"])] = link
        record["links"] = list(unique_links.values())
    return {
        "records": records,
        "invalid_rows": invalid_rows,
        "total_records": len(records),
        "total_links": sum(len(record["links"]) for record in records),
        "invalid_count": len(invalid_rows),
    }


def _parse_import_file(filename: str, content: bytes) -> dict[str, Any]:
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("服务器未安装 openpyxl，无法解析 Excel 文件") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        return _build_import_preview([list(row) for row in rows])
    if not lower_name.endswith(".txt"):
        raise ValueError("只支持 TXT、XLSX 或 XLSM 文件")
    text = content.decode("utf-8-sig", errors="replace")
    rows = []
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        delimiter = "\t" if "\t" in line else "|"
        rows.append(line.split(delimiter))
    return _build_import_preview(rows)


def preview_import(session: Session, *, filename: str, content: bytes, operator: str) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    session.query(MediaWallImportJob).filter(
        MediaWallImportJob.status == "preview",
        MediaWallImportJob.created_at < _now() - timedelta(hours=24),
    ).delete(synchronize_session=False)
    preview = _parse_import_file(filename, content)
    file_hash = hashlib.sha256(content).hexdigest()
    job = MediaWallImportJob(
        filename=_clean_text(filename, max_length=255) or "media-wall-import",
        file_hash=file_hash,
        status="preview",
        preview_json=preview,
        summary_json={key: preview[key] for key in ("total_records", "total_links", "invalid_count")},
        created_by=operator,
    )
    session.add(job)
    session.flush()
    return {"id": int(job.id), "filename": job.filename, "status": job.status, **preview}


def confirm_import(session: Session, *, job_id: int, category_ids: list[int], visibility: str, operator: str) -> dict[str, Any]:
    job = session.get(MediaWallImportJob, int(job_id))
    if job is None:
        raise LookupError("导入任务不存在")
    if job.status == "imported":
        raise ValueError("该导入任务已经确认过")
    if visibility not in MEDIA_WALL_VISIBILITIES:
        raise ValueError("无效的影视公开状态")
    validated_categories = _validate_category_ids(session, category_ids)
    preview = job.preview_json or {}
    records = preview.get("records") if isinstance(preview, dict) else []
    if not isinstance(records, list):
        raise ValueError("导入预览数据无效")
    created_count = 0
    updated_count = 0
    link_count = 0
    existing_items = session.query(MediaWallItem).all()
    items_by_title = {
        _normalize_title(item.display_title): item
        for item in existing_items
        if _normalize_title(item.display_title)
    }
    for record in records:
        title = _clean_text(record.get("display_title"), max_length=255)
        if not title:
            continue
        normalized_title = _normalize_title(title)
        item = items_by_title.get(normalized_title)
        if item is None:
            _invalidate_metadata_cache(title)
            item = MediaWallItem(
                display_title=title,
                metadata_query_title=title,
                visibility=visibility,
                sort_order=0,
                metadata_status=MEDIA_WALL_METADATA_PENDING,
                created_by=operator,
                updated_by=operator,
            )
            session.add(item)
            session.flush()
            items_by_title[normalized_title] = item
            created_count += 1
        else:
            updated_count += 1
            if str(item.metadata_status or MEDIA_WALL_METADATA_NOT_SYNCED) in {MEDIA_WALL_METADATA_NOT_SYNCED, MEDIA_WALL_METADATA_FAILED}:
                _invalidate_metadata_cache(item.metadata_query_title)
                item.metadata_status = MEDIA_WALL_METADATA_PENDING
                item.metadata_last_error = None
        if validated_categories:
            existing_ids = [int(category.id) for category in _get_categories_for_item(session, item.id)]
            _replace_item_categories(session, item_id=item.id, category_ids=[*existing_ids, *validated_categories])
        for link in record.get("links", []):
            if _upsert_link(session, item_id=item.id, url=str(link.get("url") or ""), label=None, source_type="import", is_primary=False):
                link_count += 1
    job.status = "imported"
    job.confirmed_by = operator
    job.summary_json = {
        "created_count": created_count,
        "updated_count": updated_count,
        "link_count": link_count,
    }
    job_id = int(job.id)
    summary = dict(job.summary_json)
    session.delete(job)
    session.flush()
    return {"job_id": job_id, "status": "imported", **summary}

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import timedelta
from typing import Any, Iterable

from sqlalchemy import desc, or_
from sqlalchemy.orm import Session

from app.models.models import (
    Channel,
    LearningChannelConfig,
    LearningResourceAutomationTask,
    LearningResourceAutomationTaskChannel,
    LearningResourceAutomationTaskTarget,
    LearningResourceCategory,
    LearningResourceImportJob,
    LearningResourceItem,
    LearningResourceItemCategory,
    LearningResourceLink,
    LearningResourceProcessLog,
    LearningResourceTransferJob,
    LinkTarget,
    Message,
    MessageLinkRef,
    PanTransferAccount,
    PanTransferBatch,
    PanTransferBatchItem,
    PanTransferExecutionLog,
    PanTransferPublishRecord,
    PanTransferReplacementLog,
    PanTransferSyncTask,
    PanTransferSyncTaskLog,
    ensure_runtime_storage_tables,
)
from app.services.pan_transfer.common import normalize_relative_path, sanitize_path_segment, utcnow
from app.services.pan_transfer.constants import normalize_transfer_platform
from app.services.link_check.constants import PLATFORM_QUARK
from app.services.link_check.parser import detect_platform_from_url, is_http_url, normalize_candidate_url
from app.services.pan_transfer.queue import refresh_pan_transfer_batch_summary
from app.utils.channel_utils import normalize_channel_username


LEARNING_STATUS_PENDING = "pending"
LEARNING_STATUS_TRANSFERRING = "transferring"
LEARNING_STATUS_PUBLISHED = "published"
LEARNING_STATUS_FAILED = "failed"
LEARNING_STATUS_HIDDEN = "hidden"
LEARNING_JOB_QUEUED = "queued"
LEARNING_JOB_TRANSFERRING = "transferring"
LEARNING_JOB_PUBLISHED = "published"
LEARNING_JOB_FAILED = "failed"
LEARNING_JOB_PAUSED = "paused"
LEARNING_IMPORT_MODES = {"direct", "transfer", "collect"}
MAX_IMPORT_ROWS = 5000
MAX_LEARNING_PROCESS_LOGS_PER_TASK = 5000


@dataclass(frozen=True)
class LearningAutomationContext:
    task: LearningResourceAutomationTask
    channel: LearningResourceAutomationTaskChannel
    target: LearningResourceAutomationTaskTarget
    account: PanTransferAccount
    channel_key: str
    root: str

_URL_TAG_RE = re.compile(r"^(?:https?://|www\.|https?$)", re.IGNORECASE)
_SPACE_RE = re.compile(r"\s+")
_TITLE_NOISE_RE = re.compile(r"[\[\]【】()（）{}<>《》:：,，.!！?？/\\|_+−\-·•]+")


def normalize_learning_title(value: Any) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    normalized = _TITLE_NOISE_RE.sub(" ", normalized)
    normalized = _SPACE_RE.sub("", normalized)
    return normalized[:255] or "未命名资源"


def _display_title(message: Message) -> str:
    title = str(message.title or "").strip()
    if title:
        return title[:255]
    description = str(message.description or "").strip()
    return (description.splitlines()[0].strip() if description else "未命名资源")[:255]


def _strip_wrapping_quotes(value: Any) -> str:
    text = str(value or "").replace("\ufeff", "").strip()
    quote_pairs = (("\"", "\""), ("'", "'"), ("“", "”"), ("「", "」"), ("『", "』"))
    changed = True
    while len(text) >= 2 and changed:
        changed = False
        for opening, closing in quote_pairs:
            if text.startswith(opening) and text.endswith(closing):
                text = text[len(opening):len(text) - len(closing)].strip()
                changed = True
                break
    return text


def _clean_tag(value: Any) -> str:
    return _strip_wrapping_quotes(value).lstrip("#").strip()[:120]


def classify_learning_message(message: Message) -> dict[str, Any]:
    """Return an explainable top-level category for a learning resource."""
    title = _display_title(message)
    title_lower = title.lower()

    for raw_tag in list(message.tags or []):
        tag = _clean_tag(raw_tag)
        if not tag or _URL_TAG_RE.match(tag):
            continue
        if "国考" in tag or "省考" in tag:
            return {"name": "国省考", "source": "tag", "confidence": "high"}
        if any(token in tag for token in ("人教版", "小学", "初中", "高中", "小初高")):
            return {"name": "小初高", "source": "tag", "confidence": "high"}
        if tag in {"浙江", "讲义"}:
            continue
        if any(token in tag for token in ("教师", "教资", "师范")):
            return {"name": "教师类", "source": "tag", "confidence": "high"}
        return {"name": tag, "source": "tag", "confidence": "medium"}

    title_rules = (
        (("教师资格", "教资", "教师招聘", "特岗教师", "教师编", "教招"), "教师类"),
        (("国考", "省考", "公务员", "行测", "申论"), "国省考"),
        (("考研", "研究生", "考研英语", "考研数学", "考研政治"), "考研"),
        (("小学", "初中", "高中", "中考", "高考", "人教版", "高中物理", "讲义"), "小初高"),
        (("四六级", "英语四级", "英语六级"), "英语考试"),
        (("计算机", "编程", "python", "java", "前端", "后端"), "计算机"),
    )
    for tokens, category in title_rules:
        if any(token.lower() in title_lower for token in tokens):
            return {"name": category, "source": "title_rule", "confidence": "medium"}
    return {"name": "待分类", "source": "fallback", "confidence": "low"}


def _clean_import_text(value: Any, *, max_length: int = 255) -> str:
    text = " ".join(str(value or "").replace("\u200b", " ").split())
    return _strip_wrapping_quotes(text)[:max_length].strip()


def _split_import_values(value: Any) -> list[str]:
    return [
        tag
        for part in re.split(r"[、,，|;；]+", str(value or ""))
        if (tag := _clean_tag(part))
    ]


def _is_import_header(row: list[Any]) -> bool:
    first = _clean_import_text(row[0] if row else "").lower()
    return first in {"资源标题", "标题", "名称", "title", "name"}


def _parse_learning_import_rows(rows: Iterable[list[Any]]) -> dict[str, Any]:
    materialized = [list(row) for row in rows]
    if materialized and _is_import_header(materialized[0]):
        headers = [_clean_import_text(value).lower() for value in materialized.pop(0)]
    else:
        headers = []

    category_index = next(
        (index for index, value in enumerate(headers) if value in {"分类", "类别", "category"}),
        1,
    )
    tags_index = next(
        (index for index, value in enumerate(headers) if value in {"标签", "tags", "tag"}),
        2,
    )
    reserved_indexes = {0, category_index, tags_index}
    if headers:
        link_indexes = [index for index, value in enumerate(headers) if index not in reserved_indexes and ("链接" in value or "link" in value or "url" in value)]
        if not link_indexes:
            link_indexes = [index for index in range(len(headers)) if index not in reserved_indexes]
    else:
        link_indexes = []

    records: list[dict[str, Any]] = []
    invalid_rows: list[dict[str, Any]] = []
    seen: dict[tuple[str, str], dict[str, Any]] = {}
    row_count = 0
    for row_number, row in enumerate(materialized, start=2 if headers else 1):
        if row_number > MAX_IMPORT_ROWS + (1 if headers else 0):
            invalid_rows.append({"row": row_number, "reason": f"超过 {MAX_IMPORT_ROWS} 行限制"})
            break
        if not any(_clean_import_text(value) for value in row):
            continue
        row_count += 1
        title = _clean_import_text(row[0] if row else "", max_length=255)
        category = _clean_import_text(row[category_index] if len(row) > category_index else "待分类", max_length=120) or "待分类"
        tags = _split_import_values(row[tags_index] if len(row) > tags_index else "")
        candidate_values = (
            [row[index] for index in link_indexes if index < len(row)]
            if headers
            else row[3:] if len(row) > 3 else row[1:]
        )
        links: list[dict[str, str]] = []
        for raw_link in candidate_values:
            normalized = normalize_candidate_url(str(raw_link or "").strip())
            if not normalized or not is_http_url(normalized):
                continue
            links.append({
                "url": normalized,
                "platform": detect_platform_from_url(normalized),
            })
        if not title or not links:
            invalid_rows.append({"row": row_number, "title": title, "reason": "资源标题或有效网盘链接为空"})
            continue

        key = (normalize_learning_title(title), _category_slug(category))
        record = seen.get(key)
        if record is None:
            record = {
                "display_title": title,
                "category": category,
                "tags": tags,
                "links": [],
            }
            seen[key] = record
            records.append(record)
        record["tags"] = list(dict.fromkeys([*record.get("tags", []), *tags]))
        record["links"].extend(links)

    duplicate_count = 0
    for record in records:
        unique_links: dict[str, dict[str, str]] = {}
        for link in record["links"]:
            fingerprint = hashlib.sha256(str(link["url"]).encode("utf-8")).hexdigest()
            if fingerprint in unique_links:
                duplicate_count += 1
                continue
            unique_links[fingerprint] = link
        record["links"] = list(unique_links.values())
    return {
        "records": records,
        "invalid_rows": invalid_rows,
        "row_count": row_count,
        "valid_count": len(records),
        "total_records": len(records),
        "total_links": sum(len(record["links"]) for record in records),
        "duplicate_count": duplicate_count,
        "invalid_count": len(invalid_rows),
    }


def parse_learning_import_file(*, filename: str, content: bytes) -> dict[str, Any]:
    lower_name = str(filename or "").lower()
    if lower_name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("服务器未安装 openpyxl，无法解析 Excel 文件") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        return _parse_learning_import_rows(worksheet.iter_rows(values_only=True))
    if not lower_name.endswith(".txt"):
        raise ValueError("只支持 TXT、XLSX 或 XLSM 文件")
    rows: list[list[str]] = []
    for line in content.decode("utf-8-sig", errors="replace").splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        delimiter = "\t" if "\t" in line else "|"
        rows.append(line.split(delimiter))
    return _parse_learning_import_rows(rows)


def preview_learning_import(
    session: Session,
    *,
    filename: str,
    content: bytes,
    operator: str,
) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    preview = parse_learning_import_file(filename=filename, content=content)
    file_hash = hashlib.sha256(content).hexdigest()
    job = LearningResourceImportJob(
        filename=_clean_import_text(filename, max_length=255) or "learning-resource-import",
        file_hash=file_hash,
        status="preview",
        preview_json=preview,
        summary_json={key: preview.get(key, 0) for key in ("total_records", "total_links", "invalid_count", "duplicate_count")},
        row_count=int(preview.get("row_count") or 0),
        valid_count=int(preview.get("valid_count") or 0),
        duplicate_count=int(preview.get("duplicate_count") or 0),
        error_count=int(preview.get("invalid_count") or 0),
        created_by=operator,
    )
    session.add(job)
    session.flush()
    return {"id": int(job.id), "filename": job.filename, "status": job.status, **preview}


def _get_learning_quark_account(session: Session, account_id: int | None) -> PanTransferAccount:
    account = session.get(PanTransferAccount, int(account_id)) if account_id else None
    if account is None:
        account = (
            session.query(PanTransferAccount)
            .filter(PanTransferAccount.platform == PLATFORM_QUARK, PanTransferAccount.is_enabled.is_(True))
            .order_by(PanTransferAccount.is_default.desc(), PanTransferAccount.id.asc())
            .first()
        )
    if account is None or not bool(account.is_enabled) or str(account.platform or "") != PLATFORM_QUARK:
        raise ValueError("必须选择已启用的夸克网盘账号")
    return account


def _get_or_create_import_target(session: Session, *, url: str, platform: str) -> LinkTarget:
    normalized = normalize_candidate_url(url)
    fingerprint = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    target = (
        session.query(LinkTarget)
        .filter(LinkTarget.platform == platform, LinkTarget.normalized_url_hash == fingerprint)
        .first()
    )
    if target is None:
        target = LinkTarget(
            platform=platform,
            original_url=normalized,
            normalized_url=normalized,
            normalized_url_hash=fingerprint,
            first_seen_at=utcnow(),
            last_seen_at=utcnow(),
        )
        session.add(target)
        session.flush()
    else:
        target.last_seen_at = utcnow()
        target.original_url = normalized
        session.add(target)
    return target


def _upsert_learning_import_resource(
    session: Session,
    *,
    record: dict[str, Any],
    mode: str,
    target_account_id: int | None = None,
    operator: str,
) -> tuple[LearningResourceItem, LearningResourceCategory, list[tuple[LearningResourceLink, LinkTarget | None]]]:
    title = _clean_import_text(record.get("display_title"), max_length=255) or "未命名资源"
    category_name = _clean_import_text(record.get("category"), max_length=120) or "待分类"
    category = _get_or_create_category(session, name=category_name, operator=operator)
    normalized_title = normalize_learning_title(title)
    resource = (
        session.query(LearningResourceItem)
        .filter(
            LearningResourceItem.source_channel_key == "manual-import",
            LearningResourceItem.normalized_title == normalized_title,
        )
        .first()
    )
    if resource is None:
        resource = LearningResourceItem(
            display_title=title,
            normalized_title=normalized_title,
            status=LEARNING_STATUS_PENDING,
            primary_category_id=int(category.id),
            source_channel_key="manual-import",
            latest_source_at=utcnow(),
            classification_source="import",
            classification_confidence="high",
            tags=list(record.get("tags") or []),
            created_by=operator,
            updated_by=operator,
        )
        session.add(resource)
        session.flush()
    else:
        resource.display_title = title
        resource.primary_category_id = int(category.id)
        resource.latest_source_at = utcnow()
        resource.tags = list(record.get("tags") or resource.tags or [])
        resource.updated_by = operator
        session.add(resource)
    _resource_category_link(session, resource_id=int(resource.id), category_id=int(category.id))

    links: list[tuple[LearningResourceLink, LinkTarget | None]] = []
    for raw_link in list(record.get("links") or []):
        url = normalize_candidate_url(str(raw_link.get("url") or ""))
        platform = detect_platform_from_url(url)
        if not url or not is_http_url(url):
            continue
        fingerprint = hashlib.sha256(url.encode("utf-8")).hexdigest()
        link = (
            session.query(LearningResourceLink)
            .filter(
                LearningResourceLink.resource_id == int(resource.id),
                LearningResourceLink.url_fingerprint == fingerprint,
                LearningResourceLink.target_account_id == target_account_id,
            )
            .first()
        )
        if link is None:
            link = LearningResourceLink(
                resource_id=int(resource.id),
                platform=platform,
                source_url=url,
                url_fingerprint=fingerprint,
                target_account_id=int(target_account_id) if target_account_id is not None else None,
                is_primary=False,
                is_enabled=True,
                status="pending",
            )
            session.add(link)
            session.flush()
        target = None
        if mode == "transfer":
            if platform != PLATFORM_QUARK:
                raise ValueError("提交转存目前只支持夸克网盘链接")
            target = _get_or_create_import_target(session, url=url, platform=platform)
            link.source_link_target_id = int(target.id)
        if mode == "direct":
            link.share_url = url
            link.status = "published"
            link.is_enabled = True
        elif mode == "collect":
            link.status = "pending"
            link.is_enabled = False
        else:
            link.status = "pending"
            link.is_enabled = True
        session.add(link)
        links.append((link, target))

    published_link = next((link for link, _ in links if link.status == "published" and link.share_url), None)
    if published_link is not None:
        session.query(LearningResourceLink).filter(
            LearningResourceLink.resource_id == int(resource.id),
            LearningResourceLink.id != int(published_link.id),
        ).update({"is_primary": False}, synchronize_session=False)
        published_link.is_primary = True
        resource.status = LEARNING_STATUS_PUBLISHED
        resource.last_published_at = utcnow()
    elif mode == "collect":
        resource.status = LEARNING_STATUS_PENDING
    else:
        resource.status = LEARNING_STATUS_TRANSFERRING
    session.add(resource)
    session.flush()
    return resource, category, links


def confirm_learning_import(
    session: Session,
    *,
    job_id: int,
    mode: str,
    target_account_id: int | None,
    save_root: str | None,
    operator: str,
) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    normalized_mode = str(mode or "").strip().lower()
    if normalized_mode not in LEARNING_IMPORT_MODES:
        raise ValueError("无效的导入处理模式")
    job = session.get(LearningResourceImportJob, int(job_id))
    if job is None:
        raise LookupError("导入任务不存在")
    if job.status != "preview":
        raise ValueError("该导入任务已经确认或不可用")
    account = _get_learning_quark_account(session, target_account_id) if normalized_mode == "transfer" else None
    root = normalize_relative_path(str(save_root or (account.default_save_root if account else "") or ""))
    if normalized_mode == "transfer" and not root:
        raise ValueError("提交转存前必须选择资料根目录")
    preview = dict(job.preview_json or {})
    records = preview.get("records") if isinstance(preview.get("records"), list) else []
    queued: list[tuple[LearningResourceTransferJob, LearningResourceItem, LearningResourceCategory, PanTransferAccount, LinkTarget]] = []
    created_count = 0
    link_count = 0
    for record in records:
        resource_before = session.query(LearningResourceItem).filter(
            LearningResourceItem.source_channel_key == "manual-import",
            LearningResourceItem.normalized_title == normalize_learning_title(record.get("display_title")),
        ).first()
        resource, category, links = _upsert_learning_import_resource(
            session,
            record=record,
            mode=normalized_mode,
            target_account_id=int(account.id) if account is not None else None,
            operator=operator,
        )
        if resource_before is None:
            created_count += 1
        link_count += len(links)
        if normalized_mode != "transfer" or account is None:
            continue
        for link, target in links:
            if target is None:
                continue
            transfer_job = (
                session.query(LearningResourceTransferJob)
                .filter(
                    LearningResourceTransferJob.resource_id == int(resource.id),
                    LearningResourceTransferJob.source_link_target_id == int(target.id),
                    LearningResourceTransferJob.target_account_id == int(account.id),
                )
                .first()
            )
            if transfer_job is not None and transfer_job.status in {LEARNING_JOB_QUEUED, LEARNING_JOB_TRANSFERRING}:
                continue
            if transfer_job is None:
                transfer_job = LearningResourceTransferJob(
                    resource_id=int(resource.id),
                    source_message_id=None,
                    source_link_target_id=int(target.id),
                    target_account_id=int(account.id),
                    category_id=int(category.id),
                    status=LEARNING_JOB_QUEUED,
                )
                session.add(transfer_job)
                session.flush()
            else:
                transfer_job.target_account_id = int(account.id)
                transfer_job.category_id = int(category.id)
                transfer_job.status = LEARNING_JOB_QUEUED
                transfer_job.last_error_message = None
                session.add(transfer_job)
            queued.append((transfer_job, resource, category, account, target))

    batch_id = None
    if queued:
        batch = PanTransferBatch(
            batch_type="learning_resources",
            source_scope="learning_import",
            status="running",
            created_by=operator,
            total_message_count=0,
            total_link_target_count=len(queued),
            started_at=utcnow(),
            request_json={"import_job_id": int(job.id), "target_account_id": int(account.id), "save_root": root},
            result_json={},
        )
        session.add(batch)
        session.flush()
        batch_id = int(batch.id)
        for transfer_job, resource, category, target_account, target in queued:
            item = PanTransferBatchItem(
                batch_id=int(batch.id),
                link_target_id=int(target.id),
                target_account_id=int(target_account.id),
                platform=PLATFORM_QUARK,
                short_title=resource.display_title,
                original_url=str(target.original_url or ""),
                source_message_count=0,
                source_ref_count=0,
                latest_message_title=resource.display_title,
                latest_message_time=resource.latest_source_at,
                latest_link_health="unknown",
                transfer_status="queued",
                share_status="pending",
                validation_status="pending",
                replacement_status="replaced",
                max_attempts=3,
                extra_json={
                    "learning_resource_job_id": int(transfer_job.id),
                    "skip_link_replacement": True,
                    "resolved_paths": {
                        "staging_root": normalize_relative_path("/".join(part for part in (root, category.directory_name) if part)),
                        "staging_folder_name": sanitize_path_segment(resource.display_title, fallback="未命名资源", max_length=96),
                        "resolved_path": normalize_relative_path("/".join(part for part in (root, category.directory_name, resource.display_title) if part)),
                        "transfer_layout": "independent",
                        "share_target_mode": "resource_dir",
                    },
                    "path_strategy": {"item_folder_mode": "original", "share_target_mode": "resource_dir"},
                },
            )
            session.add(item)
            session.flush()
            transfer_job.batch_id = int(batch.id)
            transfer_job.batch_item_id = int(item.id)
            resource.status = LEARNING_STATUS_TRANSFERRING
            session.add_all([transfer_job, resource])
        session.flush()
        refresh_pan_transfer_batch_summary(session, batch_id=int(batch.id))

    job.status = "queued" if queued else "completed"
    job.confirmed_by = operator
    job.summary_json = {
        "mode": normalized_mode,
        "created_count": created_count,
        "link_count": link_count,
        "queued_count": len(queued),
        "batch_id": batch_id,
    }
    session.add(job)
    session.flush()
    return {"job_id": int(job.id), "status": str(job.status), **job.summary_json}


def _category_slug(name: str) -> str:
    slug = re.sub(r"[^\w\u4e00-\u9fff]+", "-", name.strip().lower()).strip("-")
    return (slug or "uncategorized")[:160]


def _get_or_create_category(
    session: Session,
    *,
    name: str,
    operator: str | None,
) -> LearningResourceCategory:
    slug = _category_slug(name)
    category = session.query(LearningResourceCategory).filter(LearningResourceCategory.slug == slug).first()
    if category is None:
        category = LearningResourceCategory(
            name=name[:120],
            slug=slug,
            directory_name=sanitize_path_segment(name, fallback="待分类", max_length=96),
            directory_status="pending",
            created_by=operator,
            updated_by=operator,
        )
        session.add(category)
        session.flush()
    return category


def get_learning_channel_config(session: Session, channel_key: str) -> LearningChannelConfig | None:
    normalized = normalize_channel_username(channel_key)
    return (
        session.query(LearningChannelConfig)
        .filter(LearningChannelConfig.channel_key == normalized)
        .first()
    )


def is_learning_channel_enabled(session: Session, channel_key: str) -> bool:
    normalized = normalize_channel_username(channel_key)
    return bool(_get_learning_automation_contexts(session, normalized))


def _record_learning_process_log(
    session: Session,
    *,
    stage: str,
    level: str = "info",
    message: str,
    task_id: int | None = None,
    channel_id: int | None = None,
    target_id: int | None = None,
    resource_id: int | None = None,
    transfer_job_id: int | None = None,
    import_job_id: int | None = None,
    supplement_run_id: int | None = None,
    source_message_id: int | None = None,
    details: dict[str, Any] | None = None,
) -> LearningResourceProcessLog:
    row = LearningResourceProcessLog(
        task_id=task_id,
        channel_id=channel_id,
        target_id=target_id,
        resource_id=resource_id,
        transfer_job_id=transfer_job_id,
        import_job_id=import_job_id,
        supplement_run_id=supplement_run_id,
        source_message_id=source_message_id,
        stage=str(stage or "failed")[:32],
        level=str(level or "info")[:16],
        message=" ".join(str(message or "").split())[:2000],
        details_json=dict(details or {}),
    )
    session.add(row)
    session.flush()
    if task_id is not None and int(row.id) % 100 == 0:
        retention_now = utcnow()
        session.query(LearningResourceProcessLog).filter(
            LearningResourceProcessLog.task_id == int(task_id),
            LearningResourceProcessLog.created_at < retention_now - timedelta(days=90),
        ).delete(synchronize_session=False)
        session.query(LearningResourceProcessLog).filter(
            LearningResourceProcessLog.task_id == int(task_id),
            LearningResourceProcessLog.level == "info",
            LearningResourceProcessLog.stage.notin_(["failed", "retrying", "config"]),
            LearningResourceProcessLog.created_at < retention_now - timedelta(days=30),
        ).delete(synchronize_session=False)
        stale_ids = [
            int(log_id)
            for (log_id,) in session.query(LearningResourceProcessLog.id)
            .filter(LearningResourceProcessLog.task_id == int(task_id))
            .order_by(desc(LearningResourceProcessLog.id))
            .offset(MAX_LEARNING_PROCESS_LOGS_PER_TASK)
            .all()
        ]
        if stale_ids:
            session.query(LearningResourceProcessLog).filter(
                LearningResourceProcessLog.id.in_(stale_ids),
            ).delete(synchronize_session=False)
    return row


def record_learning_enqueue_failure(
    session: Session,
    *,
    message_ids: Iterable[int],
    channel_key: str,
    error_message: str | BaseException,
) -> int:
    """Persist replayable enqueue failures without aborting message ingestion."""
    normalized_ids = sorted({int(message_id) for message_id in message_ids if int(message_id) > 0})
    if not normalized_ids:
        return 0
    normalized_channel = str(channel_key or "").strip()
    try:
        normalized_channel = normalize_channel_username(normalized_channel)
    except ValueError:
        normalized_channel = normalized_channel[:255]
    contexts = _get_learning_automation_contexts(session, normalized_channel)
    channel_ids = {
        int(context.channel.channel_id)
        for context in contexts
        if context.channel.channel_id is not None
    }
    task_targets = {
        (
            int(context.task.id) if context.task.id is not None else None,
            int(context.target.id) if context.target.id is not None else None,
        )
        for context in contexts
    }
    if not task_targets:
        task_targets = {(None, None)}
    error_text = str(error_message or "学习资料自动转存入队失败").strip()[:1800]
    error_type = type(error_message).__name__ if isinstance(error_message, BaseException) else "EnqueueError"
    written = 0
    for message_id in normalized_ids:
        for task_id, target_id in task_targets:
            _record_learning_process_log(
                session,
                stage="failed",
                level="error",
                message=f"学习资料自动转存入队失败：{error_text}",
                task_id=task_id,
                channel_id=next(iter(channel_ids), None),
                target_id=target_id,
                source_message_id=message_id,
                details={
                    "replayable": True,
                    "channel_key": normalized_channel,
                    "error_type": error_type,
                },
            )
            written += 1
    return written


def _get_learning_automation_contexts(
    session: Session,
    channel_key: str,
    *,
    task_id: int | None = None,
    allow_paused_task: bool = False,
) -> list[LearningAutomationContext]:
    normalized = normalize_channel_username(channel_key)
    channel_candidates = {normalized, f"@{normalized}"}
    channels = session.query(Channel).filter(Channel.username.in_(sorted(channel_candidates))).all()
    channel_ids = {int(channel.id) for channel in channels if channel is not None}
    contexts: list[LearningAutomationContext] = []
    if channel_ids:
        task_context_query = (
            session.query(
                LearningResourceAutomationTaskChannel,
                LearningResourceAutomationTask,
                LearningResourceAutomationTaskTarget,
                PanTransferAccount,
            )
            .join(
                LearningResourceAutomationTask,
                LearningResourceAutomationTask.id == LearningResourceAutomationTaskChannel.task_id,
            )
            .join(
                LearningResourceAutomationTaskTarget,
                LearningResourceAutomationTaskTarget.task_id == LearningResourceAutomationTask.id,
            )
            .join(
                PanTransferAccount,
                PanTransferAccount.id == LearningResourceAutomationTaskTarget.target_account_id,
            )
            .filter(
                LearningResourceAutomationTaskChannel.channel_id.in_(sorted(channel_ids)),
                LearningResourceAutomationTaskChannel.is_enabled.is_(True),
                LearningResourceAutomationTask.transfer_policy == "automatic",
                LearningResourceAutomationTask.status != "deleted",
                LearningResourceAutomationTaskTarget.is_enabled.is_(True),
                PanTransferAccount.is_enabled.is_(True),
                PanTransferAccount.platform == PLATFORM_QUARK,
            )
            .order_by(
                LearningResourceAutomationTaskChannel.id.asc(),
                LearningResourceAutomationTaskTarget.id.asc(),
            )
        )
        if not allow_paused_task:
            task_context_query = task_context_query.filter(LearningResourceAutomationTask.is_enabled.is_(True))
        task_context_rows = task_context_query.all()
        for task_channel, task, target, account in task_context_rows:
            if task_id is not None and int(task.id) != int(task_id):
                continue
            root = normalize_relative_path(str(target.save_root_name or account.default_save_root or ""))
            if not root:
                continue
            contexts.append(
                LearningAutomationContext(
                    task=task,
                    channel=task_channel,
                    target=target,
                    account=account,
                    channel_key=normalized,
                    root=root,
                )
            )

    if contexts:
        return contexts

    # Keep direct callers and partially migrated installations functional until
    # the startup migration has created the task graph.
    config = get_learning_channel_config(session, normalized)
    if config is None or (not bool(config.is_enabled) and not allow_paused_task):
        return []
    legacy_task = session.query(LearningResourceAutomationTask).filter(
        LearningResourceAutomationTask.legacy_config_id == int(config.id),
    ).first()
    if legacy_task is not None and (
        str(legacy_task.status or "") == "deleted"
        or (not bool(legacy_task.is_enabled) and not allow_paused_task)
    ):
        return []
    account = _resolve_learning_account(session, config=config)
    root = normalize_relative_path(str(config.save_root or (account.default_save_root if account else "") or ""))
    if account is None or not root:
        return []
    legacy_task = LearningResourceAutomationTask(
        id=None,
        name=str(config.channel_title or normalized),
        is_enabled=True,
        transfer_policy="automatic",
        status="running",
        legacy_config_id=int(config.id),
    )
    legacy_channel = LearningResourceAutomationTaskChannel(
        id=None,
        task_id=None,
        channel_id=int(channels[0].id) if channels else None,
        is_enabled=True,
    )
    legacy_target = LearningResourceAutomationTaskTarget(
        id=None,
        task_id=None,
        target_account_id=int(account.id),
        save_root_id=f"path:{root}",
        save_root_name=root,
        is_enabled=True,
        status="ready",
    )
    return [LearningAutomationContext(legacy_task, legacy_channel, legacy_target, account, normalized, root)]


def _resolve_learning_account(
    session: Session,
    *,
    config: LearningChannelConfig,
) -> PanTransferAccount | None:
    account = session.get(PanTransferAccount, int(config.target_account_id)) if config.target_account_id else None
    if account is None:
        account = (
            session.query(PanTransferAccount)
            .filter(
                PanTransferAccount.platform == PLATFORM_QUARK,
                PanTransferAccount.is_enabled.is_(True),
            )
            .order_by(PanTransferAccount.is_default.desc(), PanTransferAccount.id.asc())
            .first()
        )
    if account is None or not bool(account.is_enabled):
        return None
    try:
        return account if normalize_transfer_platform(str(account.platform or "")) == PLATFORM_QUARK else None
    except ValueError:
        return None


def _resource_category_link(
    session: Session,
    *,
    resource_id: int,
    category_id: int,
) -> LearningResourceItemCategory:
    row = (
        session.query(LearningResourceItemCategory)
        .filter(
            LearningResourceItemCategory.resource_id == resource_id,
            LearningResourceItemCategory.category_id == category_id,
        )
        .first()
    )
    if row is None:
        row = LearningResourceItemCategory(
            resource_id=resource_id,
            category_id=category_id,
            is_primary=True,
        )
        session.add(row)
    else:
        row.is_primary = True
    session.flush()
    return row


def _build_learning_paths(
    *,
    account: PanTransferAccount,
    root: str,
    category: LearningResourceCategory,
    title: str,
) -> dict[str, str | None]:
    del account
    root = normalize_relative_path(str(root or ""))
    category_root = normalize_relative_path("/".join(part for part in (root, category.directory_name) if part))
    folder_name = sanitize_path_segment(title, fallback="未命名资源", max_length=96)
    return {
        "staging_root": category_root,
        "staging_folder_name": folder_name,
        "resolved_path": normalize_relative_path("/".join(part for part in (category_root, folder_name) if part)),
        "transfer_layout": "independent",
        "batch_folder_name": None,
        "share_target_mode": "resource_dir",
    }


def _source_fingerprint(target: LinkTarget) -> str:
    value = str(target.normalized_url_hash or target.normalized_url or target.original_url or "").strip()
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _find_quark_refs(session: Session, message_id: int) -> list[tuple[MessageLinkRef, LinkTarget]]:
    rows = (
        session.query(MessageLinkRef, LinkTarget)
        .join(LinkTarget, LinkTarget.id == MessageLinkRef.link_target_id)
        .filter(MessageLinkRef.message_id == int(message_id))
        .order_by(MessageLinkRef.link_index.asc(), MessageLinkRef.id.asc())
        .all()
    )
    result: list[tuple[MessageLinkRef, LinkTarget]] = []
    for ref, target in rows:
        try:
            if normalize_transfer_platform(str(target.platform or ref.provider_label or "")) != PLATFORM_QUARK:
                continue
        except ValueError:
            continue
        result.append((ref, target))
    return result


def _create_learning_batch(
    session: Session,
    *,
    context: LearningAutomationContext,
    jobs: list[tuple[LearningResourceTransferJob, LearningResourceItem, LearningResourceCategory, PanTransferAccount, LinkTarget, Message]],
) -> PanTransferBatch:
    batch = PanTransferBatch(
        batch_type="learning_resources",
        source_scope="learning_channel",
        status="running",
        created_by="tg-monitor",
        total_message_count=len({int(job.source_message_id) for job, *_ in jobs}),
        total_link_target_count=len(jobs),
        started_at=utcnow(),
        request_json={
            "channel_key": str(context.channel_key or ""),
            "task_id": int(context.task.id) if context.task.id is not None else None,
            "target_id": int(context.target.id) if context.target.id is not None else None,
            "target_account_id": int(jobs[0][3].id),
            "save_root": str(context.root or jobs[0][3].default_save_root or ""),
            "category_ids": sorted({int(category.id) for _, _, category, *_ in jobs}),
        },
        result_json={},
    )
    session.add(batch)
    session.flush()

    for job, resource, category, account, target, message in jobs:
        item = PanTransferBatchItem(
            batch_id=int(batch.id),
            link_target_id=int(target.id),
            target_account_id=int(account.id),
            platform=PLATFORM_QUARK,
            short_title=resource.display_title,
            original_url=str(target.original_url or ""),
            source_message_count=1,
            source_ref_count=1,
            latest_message_title=resource.display_title,
            latest_message_time=message.timestamp,
            latest_link_health="unknown",
            transfer_status="queued",
            share_status="pending",
            validation_status="pending",
            replacement_status="replaced",
            max_attempts=3,
            extra_json={
                "learning_resource_job_id": int(job.id),
                "learning_channel_key": str(context.channel_key or ""),
                "learning_task_id": int(context.task.id) if context.task.id is not None else None,
                "learning_target_id": int(context.target.id) if context.target.id is not None else None,
                "skip_link_replacement": True,
                "source_message_snapshot": {
                    "title": resource.display_title,
                    "description": str(message.description or ""),
                    "tags": list(message.tags or []),
                },
                "resolved_paths": _build_learning_paths(
                    account=account,
                    root=context.root,
                    category=category,
                    title=resource.display_title,
                ),
                "path_strategy": {
                    "transfer_layout": "independent",
                    "item_folder_mode": "original",
                    "share_target_mode": "resource_dir",
                },
            },
        )
        session.add(item)
        session.flush()
        job.batch_id = int(batch.id)
        job.batch_item_id = int(item.id)
        job.status = LEARNING_JOB_QUEUED
        resource.status = LEARNING_STATUS_TRANSFERRING
        session.add_all([job, resource])
    session.flush()
    refresh_pan_transfer_batch_summary(session, batch_id=int(batch.id))
    return batch


def enqueue_learning_resources_for_message_ids(
    session: Session,
    message_ids: Iterable[int],
    *,
    operator: str = "tg-monitor",
    task_id: int | None = None,
    supplement_run_id: int | None = None,
    allow_paused_task: bool = False,
) -> dict[str, int]:
    ensure_runtime_storage_tables()
    normalized_ids = sorted({int(message_id) for message_id in message_ids if int(message_id) > 0})
    if not normalized_ids:
        return {"messages": 0, "resources": 0, "jobs": 0, "skipped": 0, "skipped_existing": 0, "failed": 0}

    messages = session.query(Message).filter(Message.id.in_(normalized_ids)).all()
    jobs_by_context: dict[
        tuple[int, int, str],
        tuple[
            LearningAutomationContext,
            list[tuple[LearningResourceTransferJob, LearningResourceItem, LearningResourceCategory, PanTransferAccount, LinkTarget, Message]],
        ],
    ] = {}
    processed_messages = 0
    skipped = 0
    failed = 0
    skipped_existing = 0

    source_groups: dict[tuple[str, int | str], list[Message]] = {}
    for message in messages:
        source_key = (
            str(message.monitor_channel_key or "").strip(),
            int(message.monitor_message_id)
            if message.monitor_message_id is not None
            else f"db:{int(message.id)}",
        )
        source_groups.setdefault(source_key, []).append(message)
    source_position: dict[int, tuple[int, int, str | None]] = {}
    for group in source_groups.values():
        group.sort(key=lambda item: int(item.id))
        source_message_id = int(group[0].monitor_message_id) if group[0].monitor_message_id is not None else None
        for index, group_message in enumerate(group, start=1):
            source_position[int(group_message.id)] = (index, len(group), str(source_message_id) if source_message_id is not None else None)
    processed_source_keys: set[tuple[str, int | str]] = set()
    skipped_source_keys: set[tuple[str, int | str]] = set()
    logged_source_contexts: set[tuple[int, tuple[str, int | str]]] = set()

    for message in messages:
        channel_key = str(message.monitor_channel_key or "").strip()
        if not channel_key:
            skipped += 1
            continue
        try:
            channel_key = normalize_channel_username(channel_key)
        except ValueError:
            skipped += 1
            continue
        contexts = _get_learning_automation_contexts(
            session,
            channel_key,
            task_id=task_id,
            allow_paused_task=allow_paused_task,
        )
        if not contexts:
            skipped += 1
            continue

        for context in contexts:
            if context.channel.id is not None and (
                context.channel.last_processed_message_id is None
                or int(message.id) > int(context.channel.last_processed_message_id)
            ):
                context.channel.last_processed_message_id = int(message.id)
                session.add(context.channel)
            if context.channel.id is not None and message.monitor_message_id is not None and (
                context.channel.last_processed_telegram_message_id is None
                or int(message.monitor_message_id) > int(context.channel.last_processed_telegram_message_id)
            ):
                context.channel.last_processed_telegram_message_id = int(message.monitor_message_id)
                session.add(context.channel)

        quark_refs = _find_quark_refs(session, int(message.id))
        source_key = (
            channel_key,
            int(message.monitor_message_id)
            if message.monitor_message_id is not None
            else f"db:{int(message.id)}",
        )
        resource_index, resource_total, monitor_message_id = source_position.get(
            int(message.id),
            (1, 1, str(message.monitor_message_id) if message.monitor_message_id is not None else None),
        )
        title = _display_title(message)
        for context in contexts:
            context_key = (int(context.task.id or 0), source_key)
            if context_key in logged_source_contexts:
                continue
            logged_source_contexts.add(context_key)
            _record_learning_process_log(
                session,
                stage="received",
                message=f"收到原始消息：{title}",
                task_id=int(context.task.id) if context.task.id is not None else None,
                channel_id=int(context.channel.channel_id) if context.channel.channel_id is not None else None,
                supplement_run_id=supplement_run_id,
                source_message_id=int(message.id),
                details={
                    "resource_total": resource_total,
                    "monitor_message_id": monitor_message_id,
                    "raw_text_available": bool(message.monitor_raw_text),
                    "channel": str(context.channel_key or ""),
                },
            )
            _record_learning_process_log(
                session,
                stage="classified" if quark_refs else "skipped",
                message=(
                    f"解析完成：识别 {resource_total} 条资源、{len(quark_refs)} 个夸克链接"
                    if quark_refs
                    else "跳过原始消息：未识别当前任务支持的网盘链接"
                ),
                task_id=int(context.task.id) if context.task.id is not None else None,
                channel_id=int(context.channel.channel_id) if context.channel.channel_id is not None else None,
                supplement_run_id=supplement_run_id,
                source_message_id=int(message.id),
                details={
                    "resource_total": resource_total,
                    "link_count": len(quark_refs),
                    "monitor_message_id": monitor_message_id,
                    "channel": str(context.channel_key or ""),
                },
            )
        if not quark_refs:
            if source_key not in skipped_source_keys:
                skipped_source_keys.add(source_key)
                skipped += 1
            continue
        source_key = (
            channel_key,
            int(message.monitor_message_id)
            if message.monitor_message_id is not None
            else f"db:{int(message.id)}",
        )
        if source_key not in processed_source_keys:
            processed_source_keys.add(source_key)
            processed_messages += 1
        normalized_title = normalize_learning_title(title)
        classification = classify_learning_message(message)
        resource: LearningResourceItem | None = None
        source_at = message.timestamp or message.created_at or utcnow()

        for context in contexts:
            category = (
                session.get(LearningResourceCategory, int(context.channel.default_category_id))
                if context.channel.default_category_id
                else None
            )
            if category is None:
                category = _get_or_create_category(
                    session,
                    name=str(classification["name"]),
                    operator=operator,
                )
            if resource is None:
                resource = (
                    session.query(LearningResourceItem)
                    .filter(
                        LearningResourceItem.source_channel_key == channel_key,
                        LearningResourceItem.normalized_title == normalized_title,
                    )
                    .first()
                )
                if resource is None:
                    resource = LearningResourceItem(
                        display_title=title,
                        normalized_title=normalized_title,
                        status=LEARNING_STATUS_PENDING,
                        primary_category_id=int(category.id),
                        source_channel_key=channel_key,
                        first_source_message_id=int(message.id),
                        latest_source_message_id=int(message.id),
                        latest_source_at=source_at,
                        tags=[_clean_tag(tag) for tag in list(message.tags or []) if _clean_tag(tag)],
                        classification_source=str(classification["source"]),
                        classification_confidence=str(classification["confidence"]),
                        created_by=operator,
                        updated_by=operator,
                    )
                    session.add(resource)
                    session.flush()
                elif resource.latest_source_at is None or source_at >= resource.latest_source_at:
                    resource.display_title = title
                    resource.latest_source_message_id = int(message.id)
                    resource.latest_source_at = source_at
                    resource.primary_category_id = int(category.id)
                    resource.classification_source = str(classification["source"])
                    resource.classification_confidence = str(classification["confidence"])
                    resource.tags = [_clean_tag(tag) for tag in list(message.tags or []) if _clean_tag(tag)]
                    resource.updated_by = operator
                    session.add(resource)
            _resource_category_link(session, resource_id=int(resource.id), category_id=int(category.id))
            context_jobs: list[tuple[LearningResourceTransferJob, LearningResourceItem, LearningResourceCategory, PanTransferAccount, LinkTarget, Message]] = []

            for ref, target in quark_refs:
                fingerprint = _source_fingerprint(target)
                link = (
                    session.query(LearningResourceLink)
                    .filter(
                        LearningResourceLink.resource_id == int(resource.id),
                        LearningResourceLink.url_fingerprint == fingerprint,
                        LearningResourceLink.target_account_id == int(context.account.id),
                    )
                    .first()
                )
                if link is not None and link.status == "published" and link.share_url:
                    skipped_existing += 1
                    continue
                published_link = (
                    session.query(LearningResourceLink)
                    .filter(
                        LearningResourceLink.resource_id == int(resource.id),
                        LearningResourceLink.target_account_id == int(context.account.id),
                        LearningResourceLink.is_primary.is_(True),
                        LearningResourceLink.status == "published",
                        LearningResourceLink.share_url.isnot(None),
                    )
                    .first()
                )
                if link is None:
                    link = LearningResourceLink(
                        resource_id=int(resource.id),
                        platform=PLATFORM_QUARK,
                        source_url=str(target.original_url or ref.target_url or ""),
                        url_fingerprint=fingerprint,
                        target_account_id=int(context.account.id),
                        source_link_target_id=int(target.id),
                        is_primary=published_link is None,
                        is_enabled=True,
                        status="pending",
                    )
                    session.add(link)
                    session.flush()
                else:
                    link.target_account_id = int(context.account.id)
                    link.source_link_target_id = int(target.id)
                    link.is_enabled = True
                    link.status = "pending"
                    link.last_error_message = None
                    session.add(link)
                job = (
                    session.query(LearningResourceTransferJob)
                    .filter(
                        LearningResourceTransferJob.resource_id == int(resource.id),
                        LearningResourceTransferJob.source_link_target_id == int(target.id),
                        LearningResourceTransferJob.target_account_id == int(context.account.id),
                    )
                    .first()
                )
                if job is not None and job.status in {LEARNING_JOB_QUEUED, LEARNING_JOB_TRANSFERRING}:
                    skipped_existing += 1
                    continue
                if job is not None and job.status == LEARNING_JOB_PUBLISHED and link.share_url:
                    skipped_existing += 1
                    continue
                if job is not None and job.batch_item_id:
                    existing_item = session.get(PanTransferBatchItem, int(job.batch_item_id))
                    if existing_item is not None and existing_item.transfer_status in {"failed", "retry_wait", "paused"}:
                        existing_item.transfer_status = "queued"
                        existing_item.next_retry_at = None
                        existing_item.error_message = None
                        existing_item.locked_by = None
                        existing_item.locked_at = None
                        existing_item.finished_at = None
                        session.add(existing_item)
                        batch = session.get(PanTransferBatch, int(existing_item.batch_id))
                        if batch is not None:
                            batch.status = "running"
                            batch.finished_at = None
                            session.add(batch)
                        job.status = LEARNING_JOB_QUEUED
                        job.last_error_message = None
                        session.add(job)
                        _record_learning_process_log(
                            session,
                            stage="retrying",
                            message=f"重新加入转存队列：{resource.display_title}",
                            task_id=int(context.task.id) if context.task.id is not None else None,
                            channel_id=int(context.channel.channel_id) if context.channel.channel_id is not None else None,
                            target_id=int(context.target.id) if context.target.id is not None else None,
                            resource_id=int(resource.id),
                            transfer_job_id=int(job.id),
                            supplement_run_id=supplement_run_id,
                            source_message_id=int(message.id),
                            details={
                                "resource_title": str(resource.display_title or ""),
                                "category": str(category.name or ""),
                                "target_account": str(context.account.account_name or ""),
                                "channel": str(context.channel_key or ""),
                                "resource_index": resource_index,
                                "resource_total": resource_total,
                                "monitor_message_id": monitor_message_id,
                            },
                        )
                        continue
                if job is None:
                    job = LearningResourceTransferJob(
                        task_id=int(context.task.id) if context.task.id is not None else None,
                        resource_id=int(resource.id),
                        source_message_id=int(message.id),
                        source_link_target_id=int(target.id),
                        target_account_id=int(context.account.id),
                        category_id=int(category.id),
                        status=LEARNING_JOB_QUEUED,
                    )
                    session.add(job)
                    session.flush()
                else:
                    job.task_id = int(context.task.id) if context.task.id is not None else job.task_id
                    job.source_message_id = int(message.id)
                    job.target_account_id = int(context.account.id)
                    job.category_id = int(category.id)
                    job.status = LEARNING_JOB_QUEUED
                    job.last_error_message = None
                    session.add(job)
                context_jobs.append((job, resource, category, context.account, target, message))
                if published_link is None and resource.status != LEARNING_STATUS_HIDDEN:
                    resource.status = LEARNING_STATUS_TRANSFERRING
                session.add(resource)

            if context_jobs:
                group_key = (
                    int(context.task.id or context.task.legacy_config_id or 0),
                    int(context.target.id or context.account.id),
                    context.root,
                )
                group = jobs_by_context.setdefault(group_key, (context, []))
                group[1].extend(context_jobs)
                context.task.last_enqueued_at = utcnow()
                session.add(context.task)
                _record_learning_process_log(
                    session,
                    stage="enqueued",
                    message=f"已加入转存队列：资源 {resource_index}/{resource_total} · {resource.display_title}",
                    task_id=int(context.task.id) if context.task.id is not None else None,
                    channel_id=int(context.channel.channel_id) if context.channel.channel_id is not None else None,
                    target_id=int(context.target.id) if context.target.id is not None else None,
                    resource_id=int(resource.id),
                    supplement_run_id=supplement_run_id,
                    source_message_id=int(message.id),
                    details={
                        "job_count": len(context_jobs),
                        "target_account_id": int(context.account.id),
                        "target_account": str(context.account.account_name or ""),
                        "resource_title": str(resource.display_title or ""),
                        "category": str(category.name or ""),
                        "channel": str(context.channel_key or ""),
                        "resource_index": resource_index,
                        "resource_total": resource_total,
                        "monitor_message_id": monitor_message_id,
                    },
                )

    for context, jobs in jobs_by_context.values():
        if jobs:
            _create_learning_batch(session, context=context, jobs=jobs)
    return {
        "messages": processed_messages,
        "resources": len({int(job[1].id) for context, jobs in jobs_by_context.values() for job in jobs}),
        "jobs": sum(len(jobs) for context, jobs in jobs_by_context.values()),
        "skipped": skipped,
        "skipped_existing": skipped_existing,
        "failed": failed,
    }


def sync_learning_resource_after_transfer(
    session: Session,
    *,
    item: PanTransferBatchItem,
    success: bool,
    error_message: str | None = None,
) -> None:
    job_id = dict(item.extra_json or {}).get("learning_resource_job_id")
    if not job_id:
        return
    job = session.get(LearningResourceTransferJob, int(job_id))
    if job is None:
        return
    extra_json = dict(item.extra_json or {})
    task_id = int(job.task_id) if job.task_id else int(extra_json.get("learning_task_id") or 0) or None
    target_id = int(extra_json.get("learning_target_id") or 0) or None
    resource = session.get(LearningResourceItem, int(job.resource_id))
    link = (
        session.query(LearningResourceLink)
        .filter(
            LearningResourceLink.resource_id == int(job.resource_id),
            LearningResourceLink.source_link_target_id == int(job.source_link_target_id),
            LearningResourceLink.target_account_id == int(job.target_account_id),
        )
        .first()
    )
    if link is None:
        # Jobs created before target-account scoping used a nullable link target.
        link = (
            session.query(LearningResourceLink)
            .filter(
                LearningResourceLink.resource_id == int(job.resource_id),
                LearningResourceLink.source_link_target_id == int(job.source_link_target_id),
                LearningResourceLink.target_account_id.is_(None),
            )
            .first()
        )
        if link is not None:
            link.target_account_id = int(job.target_account_id)
            session.add(link)
    if success:
        job.attempt_count = max(1, int(job.attempt_count or 0) + 1)
        job.last_error_message = None
        if link is not None:
            link.share_url = str(item.new_share_url or "") or None
            link.status = "published" if link.share_url else "failed"
            link.is_enabled = bool(link.share_url)
            link.is_primary = bool(link.share_url)
            link.last_error_message = None if link.share_url else "转存完成但未生成分享链接"
            session.add(link)
            if link.share_url:
                session.query(LearningResourceLink).filter(
                    LearningResourceLink.resource_id == int(job.resource_id),
                    LearningResourceLink.id != int(link.id),
                ).update({"is_primary": False, "is_enabled": False}, synchronize_session=False)
        if resource is not None:
            if resource.status != LEARNING_STATUS_HIDDEN:
                resource.status = LEARNING_STATUS_PUBLISHED if link is not None and link.share_url else LEARNING_STATUS_FAILED
            if resource.status == LEARNING_STATUS_PUBLISHED:
                resource.last_published_at = utcnow()
            resource.updated_by = "tg-worker"
            session.add(resource)
        job.status = LEARNING_JOB_PUBLISHED if resource is None or bool(link is not None and link.share_url) else LEARNING_JOB_FAILED
        published = job.status == LEARNING_JOB_PUBLISHED
        category_name = ""
        if resource is not None and resource.primary_category_id:
            category = session.get(LearningResourceCategory, int(resource.primary_category_id))
            category_name = str(category.name or "") if category is not None else ""
        if task_id is not None:
            if published:
                _record_learning_process_log(
                    session,
                    stage="sharing",
                    message=f"分享链接已生成，准备发布：{resource.display_title if resource is not None else f'任务#{job.id}'}",
                    task_id=task_id,
                    target_id=target_id,
                    resource_id=int(resource.id) if resource is not None else None,
                    transfer_job_id=int(job.id),
                    source_message_id=int(job.source_message_id) if job.source_message_id else None,
                    details={
                        "share_url_created": True,
                        "resource_title": str(resource.display_title or "") if resource is not None else "",
                    },
                )
            _record_learning_process_log(
                session,
                stage="published" if published else "failed",
                level="info" if published else "error",
                message=(
                    f"转存完成并发布：{resource.display_title if resource is not None else f'任务#{job.id}'}"
                    if published
                    else f"转存完成但未发布：{resource.display_title if resource is not None else f'任务#{job.id}'}"
                ),
                task_id=task_id,
                target_id=target_id,
                resource_id=int(resource.id) if resource is not None else None,
                transfer_job_id=int(job.id),
                source_message_id=int(job.source_message_id) if job.source_message_id else None,
                details={
                    "share_url_created": bool(link is not None and link.share_url),
                    "resource_title": str(resource.display_title or "") if resource is not None else "",
                    "category": category_name,
                },
            )
            task = session.get(LearningResourceAutomationTask, int(task_id))
            target = session.get(LearningResourceAutomationTaskTarget, int(target_id)) if target_id else None
            if published:
                if task is not None:
                    task.last_success_at = utcnow()
                    task.last_error_message = None
                    session.add(task)
                if target is not None:
                    target.last_success_at = utcnow()
                    target.last_error_message = None
                    session.add(target)
    else:
        retry_waiting = str(item.transfer_status or "") == "retry_wait"
        job.status = LEARNING_JOB_QUEUED if retry_waiting else LEARNING_JOB_FAILED
        job.attempt_count = max(1, int(job.attempt_count or 0) + 1)
        job.last_error_message = str(error_message or "学习资料转存失败")[:2000]
        if link is not None:
            link.status = "failed"
            link.last_error_message = job.last_error_message
            session.add(link)
        if resource is not None and not retry_waiting and resource.status != LEARNING_STATUS_HIDDEN:
            fallback_link = (
                session.query(LearningResourceLink)
                .filter(
                    LearningResourceLink.resource_id == int(resource.id),
                    LearningResourceLink.id != int(link.id) if link is not None else True,
                    LearningResourceLink.is_primary.is_(True),
                    LearningResourceLink.is_enabled.is_(True),
                    LearningResourceLink.status == "published",
                    LearningResourceLink.share_url.isnot(None),
                )
                .first()
            )
            resource.status = LEARNING_STATUS_PUBLISHED if fallback_link is not None else LEARNING_STATUS_FAILED
            resource.updated_by = "tg-worker"
            session.add(resource)
        if task_id is not None:
            _record_learning_process_log(
                session,
                stage="retrying" if retry_waiting else "failed",
                level="warning" if retry_waiting else "error",
                message=(
                    f"转存失败，等待自动重试：{resource.display_title if resource is not None else f'任务#{job.id}'}"
                    if retry_waiting
                    else f"转存失败：{resource.display_title if resource is not None else f'任务#{job.id}'}"
                ),
                task_id=task_id,
                target_id=target_id,
                resource_id=int(resource.id) if resource is not None else None,
                transfer_job_id=int(job.id),
                source_message_id=int(job.source_message_id) if job.source_message_id else None,
                details={
                    "retry_waiting": retry_waiting,
                    "error": job.last_error_message,
                    "resource_title": str(resource.display_title or "") if resource is not None else "",
                },
            )
            task = session.get(LearningResourceAutomationTask, int(task_id))
            target = session.get(LearningResourceAutomationTaskTarget, int(target_id)) if target_id else None
            if task is not None:
                task.last_error_message = job.last_error_message
                session.add(task)
            if target is not None:
                target.last_error_message = job.last_error_message
                session.add(target)
    session.add(job)
    if task_id is not None:
        from app.services.learning_automation import refresh_learning_automation_task_health

        refresh_learning_automation_task_health(session, int(task_id))
    session.flush()


def retry_learning_resource(session: Session, resource_id: int) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    resource = session.get(LearningResourceItem, int(resource_id))
    if resource is None:
        raise LookupError("学习资料不存在")
    job = (
        session.query(LearningResourceTransferJob)
        .filter(LearningResourceTransferJob.resource_id == int(resource.id))
        .order_by(LearningResourceTransferJob.id.desc())
        .first()
    )
    if job is None or not job.batch_item_id:
        raise ValueError("该资源没有可重试的转存任务")
    item = session.get(PanTransferBatchItem, int(job.batch_item_id))
    if item is None:
        raise ValueError("转存任务条目不存在")
    item.transfer_status = "queued"
    item.next_retry_at = None
    item.locked_by = None
    item.locked_at = None
    item.started_at = None
    item.finished_at = None
    item.error_message = None
    item.share_status = "shared" if item.new_share_url else "pending"
    item.validation_status = "pending"
    item.replacement_status = "replaced" if dict(item.extra_json or {}).get("skip_link_replacement") else "pending"
    session.add(item)
    batch = session.get(PanTransferBatch, int(item.batch_id))
    if batch is not None:
        batch.status = "running"
        batch.finished_at = None
        batch.started_at = batch.started_at or utcnow()
        session.add(batch)
    job.status = LEARNING_JOB_QUEUED
    job.last_error_message = None
    resource.status = LEARNING_STATUS_TRANSFERRING
    session.add_all([job, resource])
    session.flush()
    refresh_pan_transfer_batch_summary(session, batch_id=int(item.batch_id))
    return {"resource_id": int(resource.id), "job_id": int(job.id), "status": str(job.status)}


def _normalize_learning_tags(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    for value in values:
        tag = _clean_tag(value)
        if tag and tag not in result:
            result.append(tag)
    return result[:50]


def update_learning_resource(
    session: Session,
    *,
    resource_id: int,
    display_title: str,
    category_name: str,
    tags: Iterable[Any],
    share_url: str,
    operator: str,
) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    resource = session.get(LearningResourceItem, int(resource_id))
    if resource is None:
        raise LookupError("学习资料不存在")

    title = _clean_import_text(display_title, max_length=255)
    if not title:
        raise ValueError("资源标题不能为空")
    normalized_title = normalize_learning_title(title)
    duplicate = (
        session.query(LearningResourceItem.id)
        .filter(
            LearningResourceItem.source_channel_key == resource.source_channel_key,
            LearningResourceItem.normalized_title == normalized_title,
            LearningResourceItem.id != int(resource.id),
        )
        .first()
    )
    if duplicate is not None:
        raise ValueError("同一来源已经存在相同标题的资源")

    category = _get_or_create_category(
        session,
        name=_clean_import_text(category_name, max_length=120) or "待分类",
        operator=operator,
    )
    resource.display_title = title
    resource.normalized_title = normalized_title
    resource.primary_category_id = int(category.id)
    resource.tags = _normalize_learning_tags(tags)
    resource.updated_by = operator
    session.add(resource)
    session.query(LearningResourceItemCategory).filter(
        LearningResourceItemCategory.resource_id == int(resource.id),
    ).delete(synchronize_session=False)
    _resource_category_link(session, resource_id=int(resource.id), category_id=int(category.id))

    normalized_share_url = normalize_candidate_url(str(share_url or ""))
    if normalized_share_url and not is_http_url(normalized_share_url):
        raise ValueError("网盘链接必须是有效的 http 或 https 地址")
    link = (
        session.query(LearningResourceLink)
        .filter(
            LearningResourceLink.resource_id == int(resource.id),
            LearningResourceLink.is_primary.is_(True),
        )
        .order_by(LearningResourceLink.id.desc())
        .first()
    )
    if normalized_share_url:
        share_fingerprint = hashlib.sha256(normalized_share_url.encode("utf-8")).hexdigest()
        matching_link = (
            session.query(LearningResourceLink)
            .filter(
                LearningResourceLink.resource_id == int(resource.id),
                LearningResourceLink.url_fingerprint == share_fingerprint,
            )
            .order_by(LearningResourceLink.id.desc())
            .first()
        )
        if matching_link is not None:
            link = matching_link
        if link is None:
            link = LearningResourceLink(
                resource_id=int(resource.id),
                platform=detect_platform_from_url(normalized_share_url),
                source_url=normalized_share_url,
                url_fingerprint=share_fingerprint,
                share_url=normalized_share_url,
                is_primary=True,
                is_enabled=True,
                status="published",
            )
            session.add(link)
            session.flush()
        else:
            link.platform = detect_platform_from_url(normalized_share_url)
            link.source_url = normalized_share_url
            link.url_fingerprint = share_fingerprint
            link.share_url = normalized_share_url
            link.is_primary = True
            link.is_enabled = True
            link.status = "published"
            link.last_error_message = None
            session.add(link)
        session.query(LearningResourceLink).filter(
            LearningResourceLink.resource_id == int(resource.id),
            LearningResourceLink.id != int(link.id) if link.id is not None else True,
        ).update({"is_primary": False}, synchronize_session=False)
        resource.status = LEARNING_STATUS_PUBLISHED
        resource.last_published_at = utcnow()
    elif link is not None:
        link.is_primary = False
        link.is_enabled = False
        link.status = "pending"
        link.last_error_message = "管理员清空了分享链接"
        session.add(link)
        resource.status = LEARNING_STATUS_PENDING
    else:
        resource.status = LEARNING_STATUS_PENDING
    session.add(resource)
    session.flush()
    result = get_admin_learning_item(session, int(resource.id))
    if result is None:
        raise LookupError("学习资料更新后无法读取")
    return result


def _delete_transfer_batch_records(
    session: Session,
    *,
    batch_item_ids: set[int],
    batch_ids: set[int],
    delete_batches: bool,
) -> dict[str, int]:
    if batch_ids:
        batch_item_ids.update(
            int(item_id)
            for (item_id,) in session.query(PanTransferBatchItem.id)
            .filter(PanTransferBatchItem.batch_id.in_(sorted(batch_ids)))
            .all()
        )
    if not batch_item_ids and not (delete_batches and batch_ids):
        return {"deleted_batches": 0}

    batch_scope_filters = []
    if batch_item_ids:
        batch_scope_filters.append(PanTransferSyncTask.source_batch_item_id.in_(sorted(batch_item_ids)))
    if batch_ids:
        batch_scope_filters.append(PanTransferSyncTask.source_batch_id.in_(sorted(batch_ids)))
    sync_task_ids = {
        int(task_id)
        for (task_id,) in session.query(PanTransferSyncTask.id)
        .filter(or_(*batch_scope_filters))
        .all()
    } if batch_scope_filters else set()
    if sync_task_ids:
        session.query(PanTransferSyncTaskLog).filter(PanTransferSyncTaskLog.task_id.in_(sorted(sync_task_ids))).delete(
            synchronize_session=False,
        )
        session.query(PanTransferSyncTask).filter(PanTransferSyncTask.id.in_(sorted(sync_task_ids))).delete(
            synchronize_session=False,
        )
    publish_scope_filters = []
    if batch_item_ids:
        publish_scope_filters.append(PanTransferPublishRecord.source_batch_item_id.in_(sorted(batch_item_ids)))
    if batch_ids:
        publish_scope_filters.append(PanTransferPublishRecord.source_batch_id.in_(sorted(batch_ids)))
    if publish_scope_filters:
        session.query(PanTransferPublishRecord).filter(or_(*publish_scope_filters)).delete(synchronize_session=False)
    if batch_item_ids:
        session.query(PanTransferReplacementLog).filter(
            PanTransferReplacementLog.batch_item_id.in_(sorted(batch_item_ids)),
        ).delete(synchronize_session=False)
        session.query(PanTransferExecutionLog).filter(
            PanTransferExecutionLog.batch_item_id.in_(sorted(batch_item_ids)),
        ).delete(synchronize_session=False)
        session.query(PanTransferBatchItem).filter(
            PanTransferBatchItem.id.in_(sorted(batch_item_ids)),
        ).delete(synchronize_session=False)

    deleted_batches = 0
    if delete_batches and batch_ids:
        deleted_batches = int(
            session.query(PanTransferBatch)
            .filter(PanTransferBatch.id.in_(sorted(batch_ids)))
            .delete(synchronize_session=False)
        )
    return {"deleted_batches": deleted_batches}


def _delete_learning_resources_by_ids(session: Session, resource_ids: Iterable[int]) -> dict[str, int]:
    normalized_ids = {int(resource_id) for resource_id in resource_ids if int(resource_id) > 0}
    if not normalized_ids:
        return {
            "deleted_resources": 0,
            "deleted_links": 0,
            "deleted_transfer_jobs": 0,
            "deleted_batches": 0,
            "deleted_categories": 0,
        }
    jobs = session.query(LearningResourceTransferJob).filter(
        LearningResourceTransferJob.resource_id.in_(sorted(normalized_ids)),
    ).all()
    batch_item_ids = {int(job.batch_item_id) for job in jobs if job.batch_item_id}
    batch_ids = {int(job.batch_id) for job in jobs if job.batch_id}
    batch_stats = _delete_transfer_batch_records(
        session,
        batch_item_ids=batch_item_ids,
        batch_ids=set(),
        delete_batches=False,
    )
    deleted_jobs = int(
        session.query(LearningResourceTransferJob)
        .filter(LearningResourceTransferJob.resource_id.in_(sorted(normalized_ids)))
        .delete(synchronize_session=False)
    )
    deleted_links = int(
        session.query(LearningResourceLink)
        .filter(LearningResourceLink.resource_id.in_(sorted(normalized_ids)))
        .delete(synchronize_session=False)
    )
    session.query(LearningResourceItemCategory).filter(
        LearningResourceItemCategory.resource_id.in_(sorted(normalized_ids)),
    ).delete(synchronize_session=False)
    deleted_resources = int(
        session.query(LearningResourceItem)
        .filter(LearningResourceItem.id.in_(sorted(normalized_ids)))
        .delete(synchronize_session=False)
    )

    empty_batch_ids = {
        int(batch_id)
        for (batch_id,) in session.query(PanTransferBatch.id)
        .filter(PanTransferBatch.id.in_(sorted(batch_ids)))
        .filter(~session.query(PanTransferBatchItem.id).filter(PanTransferBatchItem.batch_id == PanTransferBatch.id).exists())
        .all()
    }
    if empty_batch_ids:
        batch_stats["deleted_batches"] += _delete_transfer_batch_records(
            session,
            batch_item_ids=set(),
            batch_ids=empty_batch_ids,
            delete_batches=True,
        )["deleted_batches"]

    deleted_categories = 0
    for category in session.query(LearningResourceCategory).all():
        item_used = session.query(LearningResourceItemCategory.id).filter(
            LearningResourceItemCategory.category_id == int(category.id),
        ).first()
        primary_item_used = session.query(LearningResourceItem.id).filter(
            LearningResourceItem.primary_category_id == int(category.id),
        ).first()
        transfer_used = session.query(LearningResourceTransferJob.id).filter(
            LearningResourceTransferJob.category_id == int(category.id),
        ).first()
        config_used = session.query(LearningChannelConfig.id).filter(
            LearningChannelConfig.default_category_id == int(category.id),
        ).first()
        if item_used is None and primary_item_used is None and transfer_used is None and config_used is None:
            session.delete(category)
            deleted_categories += 1
    return {
        "deleted_resources": deleted_resources,
        "deleted_links": deleted_links,
        "deleted_transfer_jobs": deleted_jobs,
        "deleted_batches": int(batch_stats["deleted_batches"]),
        "deleted_categories": deleted_categories,
    }


def delete_learning_resource(session: Session, *, resource_id: int) -> dict[str, int]:
    ensure_runtime_storage_tables()
    resource = session.get(LearningResourceItem, int(resource_id))
    if resource is None:
        raise LookupError("学习资料不存在")
    result = _delete_learning_resources_by_ids(session, [int(resource.id)])
    session.flush()
    return result


def delete_all_imported_learning_resources(session: Session) -> dict[str, int]:
    ensure_runtime_storage_tables()
    resource_ids = [
        int(resource_id)
        for (resource_id,) in session.query(LearningResourceItem.id)
        .filter(LearningResourceItem.source_channel_key == "manual-import")
        .all()
    ]
    result = _delete_learning_resources_by_ids(session, resource_ids)
    import_batches = [
        batch
        for batch in session.query(PanTransferBatch)
        .filter(
            PanTransferBatch.batch_type == "learning_resources",
            PanTransferBatch.source_scope == "learning_import",
        )
        .all()
    ]
    import_batch_ids = {int(batch.id) for batch in import_batches}
    if import_batch_ids:
        result["deleted_batches"] += _delete_transfer_batch_records(
            session,
            batch_item_ids=set(),
            batch_ids=import_batch_ids,
            delete_batches=True,
        )["deleted_batches"]
    result["deleted_import_jobs"] = int(
        session.query(LearningResourceImportJob).delete(synchronize_session=False)
    )
    session.flush()
    return result


def pause_learning_channel_tasks(session: Session, channel_key: str) -> int:
    normalized = normalize_channel_username(channel_key)
    resource_ids = [
        int(value)
        for (value,) in session.query(LearningResourceItem.id)
        .filter(LearningResourceItem.source_channel_key == normalized)
        .all()
    ]
    if not resource_ids:
        return 0
    jobs = session.query(LearningResourceTransferJob).filter(
        LearningResourceTransferJob.resource_id.in_(resource_ids),
        LearningResourceTransferJob.status.in_([LEARNING_JOB_QUEUED, LEARNING_JOB_TRANSFERRING]),
    ).all()
    for job in jobs:
        if job.batch_item_id:
            item = session.get(PanTransferBatchItem, int(job.batch_item_id))
            if item is not None and item.transfer_status == "queued":
                item.transfer_status = "paused"
                session.add(item)
        job.status = LEARNING_JOB_PAUSED
        session.add(job)
    session.flush()
    return len(jobs)


def resume_learning_channel_tasks(session: Session, channel_key: str) -> int:
    normalized = normalize_channel_username(channel_key)
    resource_ids = [
        int(value)
        for (value,) in session.query(LearningResourceItem.id)
        .filter(LearningResourceItem.source_channel_key == normalized)
        .all()
    ]
    if not resource_ids:
        return 0
    jobs = session.query(LearningResourceTransferJob).filter(
        LearningResourceTransferJob.resource_id.in_(resource_ids),
        LearningResourceTransferJob.status == LEARNING_JOB_PAUSED,
    ).all()
    for job in jobs:
        if job.batch_item_id:
            item = session.get(PanTransferBatchItem, int(job.batch_item_id))
            if item is not None and item.transfer_status == "paused":
                item.transfer_status = "queued"
                session.add(item)
        job.status = LEARNING_JOB_QUEUED
        session.add(job)
    session.flush()
    return len(jobs)


def list_public_learning_categories(session: Session) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    rows = (
        session.query(LearningResourceCategory)
        .join(LearningResourceItemCategory, LearningResourceItemCategory.category_id == LearningResourceCategory.id)
        .join(LearningResourceItem, LearningResourceItem.id == LearningResourceItemCategory.resource_id)
        .filter(
            LearningResourceCategory.is_enabled.is_(True),
            LearningResourceCategory.show_in_tabs.is_(True),
            LearningResourceItem.status == LEARNING_STATUS_PUBLISHED,
        )
        .distinct()
        .order_by(LearningResourceCategory.sort_order.asc(), LearningResourceCategory.id.asc())
        .all()
    )
    return [
        {
            "id": int(row.id),
            "name": _clean_import_text(row.name, max_length=120),
            "slug": str(row.slug or ""),
            "sort_order": int(row.sort_order or 0),
            "resource_count": int(
                session.query(LearningResourceItemCategory)
                .join(LearningResourceItem, LearningResourceItem.id == LearningResourceItemCategory.resource_id)
                .filter(
                    LearningResourceItemCategory.category_id == int(row.id),
                    LearningResourceItem.status == LEARNING_STATUS_PUBLISHED,
                )
                .count()
            ),
        }
        for row in rows
    ]


def list_public_learning_resources(
    session: Session,
    *,
    category_slug: str | None = None,
    search: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    query = (
        session.query(LearningResourceItem, LearningResourceCategory, LearningResourceLink, Message)
        .join(LearningResourceCategory, LearningResourceCategory.id == LearningResourceItem.primary_category_id, isouter=True)
        .join(
            LearningResourceLink,
            (LearningResourceLink.resource_id == LearningResourceItem.id)
            & LearningResourceLink.is_primary.is_(True)
            & LearningResourceLink.is_enabled.is_(True)
            & (LearningResourceLink.status == "published")
            & LearningResourceLink.share_url.isnot(None),
        )
        .join(Message, Message.id == LearningResourceItem.latest_source_message_id, isouter=True)
        .filter(LearningResourceItem.status == LEARNING_STATUS_PUBLISHED)
    )
    if category_slug and category_slug not in {"all", ""}:
        query = query.filter(LearningResourceCategory.slug == str(category_slug).strip())
    if search and search.strip():
        keyword = f"%{search.strip()}%"
        query = query.filter(
            or_(LearningResourceItem.display_title.ilike(keyword), LearningResourceItem.normalized_title.ilike(keyword))
        )
    rows = query.order_by(
        LearningResourceItem.latest_source_at.desc().nullslast(),
        LearningResourceItem.updated_at.desc(),
        LearningResourceItem.id.desc(),
    ).offset(max(0, int(offset or 0))).limit(max(1, min(int(limit or 50), 100))).all()
    return [
        {
            "id": int(item.id),
            "display_title": _clean_import_text(item.display_title, max_length=255),
            "category": {
                "id": int(category.id),
                "name": _clean_import_text(category.name, max_length=120),
                "slug": str(category.slug or ""),
            }
            if category is not None
            else None,
            "tags": [
                tag
                for raw_tag in list((message.tags if message is not None else None) or item.tags or [])
                if (tag := _clean_tag(raw_tag))
            ],
            "share_url": str(link.share_url or ""),
            "platform": str(link.platform or PLATFORM_QUARK),
            "updated_at": item.latest_source_at or item.updated_at,
            "created_at": item.created_at,
        }
        for item, category, link, message in rows
    ]


def count_public_learning_resources(
    session: Session,
    *,
    category_slug: str | None = None,
    search: str | None = None,
) -> int:
    ensure_runtime_storage_tables()
    query = (
        session.query(LearningResourceItem.id)
        .join(LearningResourceCategory, LearningResourceCategory.id == LearningResourceItem.primary_category_id, isouter=True)
        .join(
            LearningResourceLink,
            (LearningResourceLink.resource_id == LearningResourceItem.id)
            & LearningResourceLink.is_primary.is_(True)
            & LearningResourceLink.is_enabled.is_(True)
            & (LearningResourceLink.status == "published")
            & LearningResourceLink.share_url.isnot(None),
        )
        .filter(LearningResourceItem.status == LEARNING_STATUS_PUBLISHED)
    )
    if category_slug and category_slug not in {"all", ""}:
        query = query.filter(LearningResourceCategory.slug == str(category_slug).strip())
    if search and search.strip():
        keyword = f"%{search.strip()}%"
        query = query.filter(
            or_(LearningResourceItem.display_title.ilike(keyword), LearningResourceItem.normalized_title.ilike(keyword))
        )
    return int(query.count())


def _serialize_admin_learning_item(
    item: LearningResourceItem,
    category: LearningResourceCategory | None,
    link: LearningResourceLink | None,
    message: Message | None,
) -> dict[str, Any]:
    return {
        "id": int(item.id),
        "display_title": _clean_import_text(item.display_title, max_length=255),
        "status": str(item.status or ""),
            "category": {"id": int(category.id), "name": _clean_import_text(category.name, max_length=120), "slug": str(category.slug or "")} if category else None,
        "tags": [
            tag
            for raw_tag in list((message.tags if message is not None else None) or item.tags or [])
            if (tag := _clean_tag(raw_tag))
        ],
        "share_url": str(link.share_url or "") if link else None,
        "platform": str(link.platform or PLATFORM_QUARK) if link else "",
        "source_channel_key": str(item.source_channel_key or ""),
        "latest_source_at": item.latest_source_at,
        "updated_at": item.updated_at,
    }


def list_admin_learning_items(
    session: Session,
    *,
    search: str | None = None,
    category_slug: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> list[dict[str, Any]]:
    ensure_runtime_storage_tables()
    query = session.query(LearningResourceItem, LearningResourceCategory, LearningResourceLink, Message).join(
        LearningResourceCategory,
        LearningResourceCategory.id == LearningResourceItem.primary_category_id,
        isouter=True,
    ).join(
        LearningResourceLink,
        (LearningResourceLink.resource_id == LearningResourceItem.id)
        & LearningResourceLink.is_primary.is_(True),
        isouter=True,
    ).join(
        Message,
        Message.id == LearningResourceItem.latest_source_message_id,
        isouter=True,
    ).order_by(
        LearningResourceItem.latest_source_at.desc().nullslast(),
        LearningResourceItem.id.desc(),
    )
    if search and search.strip():
        keyword = f"%{search.strip()}%"
        query = query.filter(
            or_(LearningResourceItem.display_title.ilike(keyword), LearningResourceItem.normalized_title.ilike(keyword))
        )
    if category_slug and category_slug not in {"all", ""}:
        query = query.filter(LearningResourceCategory.slug == str(category_slug).strip())
    rows = query.offset(max(0, int(offset or 0))).limit(max(1, min(int(limit or 50), 100))).all()
    return [_serialize_admin_learning_item(item, category, link, message) for item, category, link, message in rows]


def count_admin_learning_items(
    session: Session,
    *,
    search: str | None = None,
    category_slug: str | None = None,
) -> int:
    ensure_runtime_storage_tables()
    query = session.query(LearningResourceItem.id).join(
        LearningResourceCategory,
        LearningResourceCategory.id == LearningResourceItem.primary_category_id,
        isouter=True,
    )
    if search and search.strip():
        keyword = f"%{search.strip()}%"
        query = query.filter(
            or_(LearningResourceItem.display_title.ilike(keyword), LearningResourceItem.normalized_title.ilike(keyword))
        )
    if category_slug and category_slug not in {"all", ""}:
        query = query.filter(LearningResourceCategory.slug == str(category_slug).strip())
    return int(query.count())


def get_admin_learning_item(session: Session, resource_id: int) -> dict[str, Any] | None:
    row = (
        session.query(LearningResourceItem, LearningResourceCategory, LearningResourceLink, Message)
        .join(LearningResourceCategory, LearningResourceCategory.id == LearningResourceItem.primary_category_id, isouter=True)
        .join(
            LearningResourceLink,
            (LearningResourceLink.resource_id == LearningResourceItem.id)
            & LearningResourceLink.is_primary.is_(True),
            isouter=True,
        )
        .join(Message, Message.id == LearningResourceItem.latest_source_message_id, isouter=True)
        .filter(LearningResourceItem.id == int(resource_id))
        .first()
    )
    if row is None:
        return None
    return _serialize_admin_learning_item(*row)


def list_learning_admin_overview(session: Session) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    configs = {str(row.channel_key): row for row in session.query(LearningChannelConfig).all()}
    channels = []
    for channel in session.query(Channel).order_by(Channel.id.asc()).all():
        try:
            key = normalize_channel_username(channel.username)
        except ValueError:
            key = str(channel.username or "").strip()
        if not key:
            continue
        config = configs.get(key)
        channels.append({
            "channel_id": int(channel.id),
            "channel_key": key,
            "channel_title": str((config.channel_title if config else None) or key),
            "config_id": int(config.id) if config else None,
            "is_enabled": bool(config.is_enabled) if config else False,
            "target_account_id": int(config.target_account_id) if config and config.target_account_id else None,
            "save_root": str(config.save_root or "") if config else "",
            "backfill_policy": str(config.backfill_policy or "new_messages_only") if config else "new_messages_only",
            "last_processed_message_id": int(config.last_processed_message_id) if config and config.last_processed_message_id else None,
            "last_error_message": str(config.last_error_message or "") if config else None,
        })
    accounts = [
        {
            "id": int(row.id),
            "platform": str(row.platform or ""),
            "account_name": str(row.account_name or ""),
            "default_save_root": str(row.default_save_root or ""),
            "is_enabled": bool(row.is_enabled),
        }
        for row in session.query(PanTransferAccount)
        .filter(PanTransferAccount.platform == PLATFORM_QUARK)
        .order_by(PanTransferAccount.is_default.desc(), PanTransferAccount.id.asc())
        .all()
    ]
    categories = [
        {
            "id": int(row.id),
            "name": _clean_import_text(row.name, max_length=120),
            "slug": str(row.slug or ""),
            "directory_name": str(row.directory_name or ""),
            "is_enabled": bool(row.is_enabled),
            "show_in_tabs": bool(row.show_in_tabs),
            "sort_order": int(row.sort_order or 0),
            "resource_count": int(
                session.query(LearningResourceItemCategory)
                .filter(LearningResourceItemCategory.category_id == int(row.id))
                .count()
            ),
        }
        for row in session.query(LearningResourceCategory)
        .order_by(LearningResourceCategory.sort_order.asc(), LearningResourceCategory.id.asc())
        .all()
    ]
    return {"channels": channels, "accounts": accounts, "categories": categories}


def update_learning_channel_config(
    session: Session,
    *,
    channel_id: int,
    payload: dict[str, Any],
    operator: str,
) -> dict[str, Any]:
    ensure_runtime_storage_tables()
    channel = session.get(Channel, int(channel_id))
    if channel is None:
        raise LookupError("监听频道不存在")
    channel_key = normalize_channel_username(channel.username)
    config = get_learning_channel_config(session, channel_key)
    if config is None:
        config = LearningChannelConfig(channel_key=channel_key, created_by=operator)
        session.add(config)
        session.flush()
    account_id = payload.get("target_account_id")
    if account_id is not None:
        account = session.get(PanTransferAccount, int(account_id))
        if account is None or not bool(account.is_enabled) or str(account.platform or "") != PLATFORM_QUARK:
            raise ValueError("必须选择已启用的夸克网盘账号")
        config.target_account_id = int(account.id)
    elif config.target_account_id is None:
        account = _resolve_learning_account(session, config=config)
        if account is not None:
            config.target_account_id = int(account.id)
    if "save_root" in payload and payload.get("save_root") is not None:
        config.save_root = normalize_relative_path(str(payload.get("save_root") or ""))
    if "is_enabled" in payload:
        config.is_enabled = bool(payload.get("is_enabled"))
    if config.is_enabled:
        account = _resolve_learning_account(session, config=config)
        if account is None:
            raise ValueError("开启前必须配置已启用的夸克网盘账号")
        if not normalize_relative_path(str(config.save_root or account.default_save_root or "")):
            raise ValueError("开启前必须配置保存根目录")
    config.channel_title = str(payload.get("channel_title") or channel.username)[:255]
    config.backfill_policy = "new_messages_only"
    config.transfer_policy = "quark_only"
    config.auto_create_category_dirs = True
    config.updated_by = operator
    config.last_error_message = None
    session.add(config)
    from app.services.learning_automation import sync_legacy_learning_channel_config

    sync_legacy_learning_channel_config(
        session,
        config=config,
        channel_id=int(channel.id),
        operator=operator,
    )
    if not config.is_enabled:
        pause_learning_channel_tasks(session, channel_key)
    else:
        resume_learning_channel_tasks(session, channel_key)
    session.flush()
    return {
        "channel_id": int(channel.id),
        "channel_key": channel_key,
        "config_id": int(config.id),
        "is_enabled": bool(config.is_enabled),
        "target_account_id": int(config.target_account_id) if config.target_account_id else None,
        "save_root": str(config.save_root or ""),
    }


__all__ = [
    "classify_learning_message",
    "count_admin_learning_items",
    "count_public_learning_resources",
    "confirm_learning_import",
    "delete_all_imported_learning_resources",
    "delete_learning_resource",
    "enqueue_learning_resources_for_message_ids",
    "get_admin_learning_item",
    "get_learning_channel_config",
    "is_learning_channel_enabled",
    "list_admin_learning_items",
    "list_learning_admin_overview",
    "list_public_learning_categories",
    "list_public_learning_resources",
    "pause_learning_channel_tasks",
    "parse_learning_import_file",
    "preview_learning_import",
    "record_learning_enqueue_failure",
    "resume_learning_channel_tasks",
    "retry_learning_resource",
    "sync_learning_resource_after_transfer",
    "update_learning_resource",
    "update_learning_channel_config",
]
